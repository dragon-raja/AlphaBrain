"""Read-only research evidence -> editable biweekly report and same-content PDF.

No policy execution, training, or evaluation-controller modification.
Dependencies are isolated in /tmp, not installed into the experiment environment.
"""
from pathlib import Path
import sys, json, hashlib, html
sys.path[:0] = ['/tmp/dsol-weekly-docx-deps-20260909',
                '/root/.cache/uv/archive-v0/XO8_y4zBEt33DOY7wfN64']
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, fontManager
import pymupdf as fitz
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent
ASSETS = OUT / '_assets'
ASSETS.mkdir(exist_ok=True)
REPO = Path('/workspace/projects/alphabrain-dsol-paper1')
EXP = Path('/share/longjunyu/alphabrain/experiments')
NEW = EXP / 'dsol-standard-initial-aa0-v1-20260908'
OLD = EXP / 'dsol-view-landscape-v1-20260908'
STEM = '科研进展双周报_20260827-20260909'
font = FontProperties(fname='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
plt.rcParams.update({'font.family': font.get_name(), 'axes.unicode_minus': False,
                     'font.size': 11, 'axes.spines.top': False, 'axes.spines.right': False})

sources = [
    REPO/'docs/dsol_paper1/training_match_execution_status_20260907_zh.md',
    REPO/'docs/dsol_paper1/libero_expanded_a_constructed_gate_20260827_zh.md',
    REPO/'docs/dsol_paper1/view_repeatability_root_cause_20260908_zh.md',
    REPO/'docs/dsol_paper1/render_protocol_bridge_and_historical_validity_20260908_zh.md',
    REPO/'docs/dsol_paper1/standard_initial_aa0_execution_20260908_zh.md',
    OLD/'checkpoint-audit/completed_run_receipt.json',
    OLD/'metric-oracle-comparison-v2/broad_summary.json',
    OLD/'render-protocol-bridge-v1/analysis-recovered-v1.json',
    NEW/'release.json', NEW/'assets/aa0-a/00/policy_inputs.npz',
]

rules = json.loads(sources[6].read_text())['overall']['rules']
keys = ['canonical', 'min_accel', 'accel_top10_max_visibility', 'statewise_empirical_max_O32']
vals = [rules[k]['success']['mean']*100 for k in keys]
assert np.allclose(vals, [64.697265625,61.376953125,60.498046875,75.87890625])

# Historical spatial distributions are auxiliary illustrations, not new outcomes.
sources += [OLD/'broad-existing/verified_arrays.npz',
            OLD/'broad-existing/candidate_geometry.json']
with np.load(sources[-2]) as a:
    success = a['success'].mean((0,2))
    accel = a['accel'].mean((0,2))
    visibility = a['visibility'].mean(0)
    candidate_ids = a['candidate_ids'].tolist()
poses_by_id = {x['pose_id']:x for x in json.loads(sources[-1].read_text())['poses']}
poses = [poses_by_id[x] for x in candidate_ids]
xyz=np.array([[x[k] for k in ['azimuth_deg','elevation_deg','radius_scale']] for x in poses])
fig=plt.figure(figsize=(8,3.1))
for j,(v,title,cmap) in enumerate([
    (success*100,'续跑成功率（%）','viridis'),
    (accel,'平均 Accel','magma'),
    (visibility*100,'任务实体可见像素占比（%）','cividis')]):
    ax=fig.add_subplot(1,3,j+1,projection='3d')
    low,high=(0,100) if j==0 else (float(v.min()),float(v.max()))
    sc=ax.scatter(*xyz[1:].T,c=v[1:],s=16,cmap=cmap,vmin=low,vmax=high,depthshade=False)
    ax.scatter(*xyz[0],c=[v[0]],s=65,marker='*',edgecolor='black',
               cmap=cmap,vmin=low,vmax=high,depthshade=False)
    ax.set_title(title,fontsize=9,pad=4)
    ax.set_xlabel('方位角偏移',fontsize=7,labelpad=-3)
    ax.set_ylabel('仰角偏移',fontsize=7,labelpad=-3)
    ax.set_zlabel('半径倍率',fontsize=7,labelpad=-7)
    ax.tick_params(labelsize=6,pad=-1);ax.view_init(22,-58)
    cb=fig.colorbar(sc,ax=ax,orientation='horizontal',shrink=.72,pad=.13,fraction=.04)
    cb.ax.tick_params(labelsize=6)
fig.subplots_adjust(left=.01,right=.98,bottom=.13,top=.92,wspace=.13)
fig.savefig(ASSETS/'historical_spatial_maps.png',dpi=220);plt.close(fig)

# Freeze a progress snapshot; intentionally do not read held-out outcome labels.
snapshot_file = OUT/'_progress_snapshot.json'
if snapshot_file.exists():
    snapshot=json.loads(snapshot_file.read_text())
else:
    from datetime import datetime, timezone
    snapshot={'observed_at_utc':datetime.now(timezone.utc).isoformat(),'hosts':{}}
    for host in ['fresh','gnu']:
        d=NEW/'hosts'/host
        snapshot['hosts'][host]={
            'status':json.loads((d/'status.json').read_text()),
            'first_wave':json.loads((d/'dense-init-0-complete.json').read_text()),
            'initial_gate':json.loads((d/'gate-pass.json').read_text()),
            'dynamic_gate':json.loads((NEW/'scheduling/dynamic-v1'/host/'cross-host-gate-pass.json').read_text()),
        }
    snapshot_file.write_text(json.dumps(snapshot,ensure_ascii=False,indent=2)+'\n')
assert all(v['first_wave']['episodes']==24832 for v in snapshot['hosts'].values())
from datetime import datetime, timedelta
asof = (datetime.fromisoformat(snapshot['observed_at_utc'])+timedelta(hours=8)).strftime('%m 月 %d 日 %H:%M')

pages=[]
def page():
    x=[];pages.append(x);return x
def add(p,kind,text='',**kw):p.append(dict(kind=kind,text=text,**kw))
def table(p,heads,rows,widths=None):add(p,'table',heads=heads,rows=rows,widths=widths)



p=page()
add(p,'title','科研进展周报')
add(p,'meta','汇报人：龙君昱    学号：26B951131')
add(p,'meta','汇报周期：2026 年 8 月 27 日—9 月 9 日')
add(p,'h1','一、这两周的主要工作')
add(p,'p','这两周围绕“训练如何影响视角泛化，以及怎样选择有利于任务执行的视角”开展实验。既考察相机扰动后能否找到模型熟悉的视角、恢复性能，也考察多视角训练后是否仍有值得选择的候选空间。主要工作包括以下四项。')
add(p,'h2','1. 对比单视角与多视角训练，分析训练覆盖的影响')
add(p,'p','为了区分视角覆盖与其他训练设置的影响，补训了与 Broad64 条件匹配的规范视角模型。两组保留相同的腕部图像、语言和动作监督，主要区别是外部训练图像是否覆盖多个位姿。目前模型训练已完成，正在相同任务和初态下比较成功率分布。')
add(p,'p','这一对照主要看：多视角训练是否扩大了可成功执行的视角范围，以及训练覆盖增加后，选择视角的收益是否随之变化。')
add(p,'h2','2. 评估候选视角收益，区分固定摆放与按初态选择')
add(p,'p','之前某些视角在单次评测中成功，并不能说明它们稳定更好。因此整理了完整候选评测，在多套推理噪声下比较规范视角、每任务最佳固定视角和逐起点 Oracle，检查收益来自任务差异，还是同一任务的不同初态需要不同视角。')
add(p,'p','已有恢复状态实验已完成整理。为检验这些发现能否适用于完整任务，新一轮改为从标准初态开始，选定外部视角后保持固定，评测整个任务的成功率，目前仍在进行。')
add(p,'h2','3. 分析 Accel 等指标能否识别有效视角')
add(p,'p','候选中存在更好的视角，不代表实际能够选出来。为此整理了成功率、Accel 和可见性的空间分布，并比较最低 Accel、可见性及组合规则的选取效果。重点是观察低 Accel 区域与高成功率区域是否对应，为后续判断指标能否识别模型熟悉的视角提供依据。')
add(p,'p','这里不只要求方法超过规范视角。如果指标能在相机扰动后选回较适应的区域，使成功率高于保持扰动视角，也是一种有效的视角利用方式。两种训练模型的分布对照将用于进一步检验这一点。')
add(p,'h2','4. 检查可见性增加是否带来额外收益')
add(p,'p','为了区分“视角更熟悉”和“任务信息更多”，整理了信息视角、规范视角及等幅移动对照的实验，并检查补充信息视角训练后的变化。这部分用于判断可见像素增加是否真正帮助任务完成，以及可见性是否适合作为 Accel 之外的补充指标。')
add(p,'p','同时修复了渲染重复性问题，统一候选图像打分与闭环评测的输入。后续将用修复后的评测重新检查上述关系。')

p=page()
add(p,'h1','二、已有结果与目前的判断')
add(p,'p','先用已有 Broad64 数据整理了下面的空间图和结果表。这批实验包含 8 个任务、64 个恢复起点，部分起点位于任务中途；新评测则从标准初态完成整个任务。两批结果需要分别看。')
add(p,'h2','1. Accel 与成功率的空间分布')
add(p,'image',file='historical_spatial_maps.png',width=17.0)
add(p,'caption','图 1  成功率、平均 Accel 和任务实体可见性。每个点对应一个候选视角，颜色为 64 个起点的平均值；三个坐标为方位角偏移、仰角偏移和半径倍率，星号为规范视角。')
add(p,'p','总体分布还不能说明低 Accel 区域就对应高成功率区域，也不能直接解释为模型熟悉的区域。需要结合下面的规则对照，并进一步查看任务和初态之间的差异。Accel 衡量的是 Flow 去噪轨迹变化，不是机械臂运动加速度。')
add(p,'h2','2. 候选 Oracle 与实际选择规则')
table(p,['视角选择方式','恢复后续跑成功率'],[
 ['97 个候选均匀选择的平均表现','58.31%'],
 ['最低平均 Accel','61.38%'],
 ['Accel 前 10 名中按可见性选择','60.50%'],
 ['规范视角','64.70%'],
 ['逐起点经验 Oracle','75.88%'],
],[.72,.28])
add(p,'p','这里的 Oracle 是对每个起点，先计算各视角在 32 套噪声下的平均成功率，再取最高值；不同起点可以选择不同视角，不是每次噪声试验后另选赢家。')
add(p,'p','旧数据中，Oracle 比规范视角高 11.18 个百分点，但最低 Accel 和组合规则没有选出同样的收益。这个差距值得继续研究，不过 75.88% 是有限样本上选出的经验成绩。修复评测后，需要重新确认候选收益是否保留，以及指标与成功率的关系有没有变化。')
add(p,'p','此前可见性实验也提醒我，目标看得更多不一定更容易成功：两个任务、12 个恢复状态中，信息视角为 75.0%，规范视角为 83.3%。因此，后面会把训练视角距离、可见性和任务难度一起分析，而不只依靠像素占比选视角。')

p=page()
add(p,'h1','三、后续工作')
add(p,'p','先完成正在运行的评测，再用同一批数据分析熟悉视角、候选 Oracle 和指标选择。暂时不把方法限定为学习一个选择器，先看哪些规律能够稳定出现。')
table(p,['下一步','具体分析'],[
 ['1. 重画空间分布','对两个模型分别画成功率、平均 Accel、Accel 跨噪声波动和可见性的热力图或三维分布图。同一任务、同一初态并排比较，再看多视角训练主要改善了哪些区域。'],
 ['2. 检查候选收益','比较规范视角、每任务最佳固定视角和逐初态经验 Oracle。看收益主要来自任务之间的差异，还是同一任务不同初态确实需要不同视角；同时比较两种训练条件。'],
 ['3. 验证熟悉视角恢复','预先选定扰动视角，比较保持扰动、返回规范视角、选择最近训练位姿，以及按 Accel 和可见性选择。主要看相对扰动视角恢复了多少成功率，再看是否优于简单返回规范位置。'],
 ['4. 分析指标有效的条件','将候选到训练视角的距离、图像相似度、目标可见性与成功率对应起来。重点看低 Accel 是否集中在熟悉区域，以及几何距离相近时，可见性差异是否还能带来收益。'],
 ['5. 再决定选择方法','如果单一指标不够，但多个指标之间有互补关系，再尝试组合打分或轻量回归、排序模型。用留出的初态检验，既与简单规则比较，也保留 Oracle 作为参照。'],
],[.25,.75])
add(p,'h2','复用已有评测数据')
add(p,'p','空间图、Oracle、固定指标规则的成功率，都能从完整评测矩阵中计算，不需要为每种分析重跑 97 个视角。新的图像相似度指标通常只需补算特征；如果最终方法需要验证新噪声下的稳定性，再补少量独立评测。')
add(p,'p','选择规则和参数只用初态 0、1 开发，初态 2、3 留作测试。同一初态的不同候选不会拆到两边，避免选择器在同一场景上既训练又测试。')
add(p,'h2','目前先不扩展的部分')
add(p,'p','现在先验证任务开始前选定固定视角的效果，还没有实际执行相机移动。计算全部候选图的指标，也不代表实体相机可以直接看到这些视角。如果这批实验能找到稳定的选择依据，再补相机搜索、移动代价或少量真机实验。RoboCasa 和更多初态也放在这一轮分析之后安排。')
add(p,'p','我希望这轮最终能说清楚两件事：指标是否能帮助模型从不利视角恢复任务能力；多视角训练后是否仍有值得选择的候选空间，以及什么信息能帮助我们找到它。')

def esc(x):return html.escape(str(x)).replace('\n','<br/>')
def html_block(b):
    k=b['kind'];t=esc(b.get('text',''))
    if k=='meta':t=t.replace('26B951131','<span style="font-family:latin">26B951131</span>')
    if k=='table':
        ws=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
        head='<tr>'+''.join(f'<th style="width:{w*100:.1f}%">{esc(v)}</th>' for v,w in zip(b['heads'],ws))+'</tr>'
        rows=''.join('<tr>'+''.join(f'<td>{esc(v)}</td>' for v in row)+'</tr>' for row in b['rows'])
        return '<table>'+head+rows+'</table>'
    if k=='image':return f'<div class="image"><img src="{b["file"]}" width="100%"/></div>'
    if k in ['h1','h2','h3']:return f'<{k}>{t}</{k}>'
    return f'<p class="{k}">{t}</p>'

css='''
@font-face {font-family:report;src:url(NotoSansCJK-Regular.ttc);}
@font-face {font-family:report;src:url(NotoSansCJK-Bold.ttc);font-weight:bold;}
@font-face {font-family:latin;src:url(DejaVuSans.ttf);}
body {font-family:report;font-size:10.5pt;line-height:1.48;color:#171717;}
p {margin:0 0 7pt 0;} .title{font-size:20pt;font-weight:bold;margin-bottom:13pt;}
.meta {font-size:10pt;margin-bottom:3pt;} h1{font-size:15pt;margin:13pt 0 8pt;}
h2{font-size:11.5pt;margin:10pt 0 6pt;} h3{font-size:11.5pt;margin:9pt 0 7pt;}
table{width:100%;border-collapse:collapse;margin:5pt 0 9pt;table-layout:fixed;}
th,td{border:0.55pt solid #cdd3d7;padding:5pt 6pt;vertical-align:top;font-size:9.6pt;line-height:1.4;}
th{background:#edf0f2;text-align:left;font-weight:bold;}
.caption{font-size:9pt;color:#555;line-height:1.4;margin:2pt 0 9pt;}
.note{font-size:10pt;background:#f2f3f4;padding:8pt;line-height:1.45;margin-top:7pt;}
.small{font-size:8.5pt;color:#666;margin-top:9pt;}.image{margin:4pt 0;}
'''
pdf=fitz.open();archive=fitz.Archive([str(ASSETS),'/usr/share/fonts/opentype/noto','/usr/share/fonts/truetype/dejavu'])
qa=[]
for i,blocks in enumerate(pages):
    pg=pdf.new_page(width=595.28,height=841.89)
    body=''.join(html_block(b) for b in blocks)
    spare,scale=pg.insert_htmlbox(fitz.Rect(54,46,541.28,791),body,css=css,archive=archive,scale_low=1)
    if spare < 0:raise RuntimeError(f'Page {i+1} overflow; shorten content, do not auto shrink')
    pg.insert_htmlbox(fitz.Rect(54,806,470,830),'科研进展周报 · VLA 视角利用 · 2026.08.27—09.09',css=css+'body{font-size:8pt;color:#777;}',archive=archive)
    pg.insert_text((527,818),str(i+1),fontsize=8,color=(.45,.45,.45))
    qa.append({'page':i+1,'spare_height':spare,'scale':scale})
pdf.set_metadata({'title':STEM,'author':'龙君昱','subject':'VLA 视角适应性与评价指标研究'})
pdf.save(OUT/(STEM+'.pdf'),garbage=4,deflate=True)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Heading 1',15),('Heading 2',11.5),('Heading 3',11.5)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True
    st.font.color.rgb=RGBColor.from_string('171717');st.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0);st.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0]
footer.text='科研进展周报 · VLA 视角利用\t'
footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.8),WD_ALIGN_PARAGRAPH.RIGHT)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('777777')
def fmt_runs(par,size=10.5,bold=False,color='171717'):
    for r in par.runs:
        r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')

for pi,blocks in enumerate(pages):
    if pi:doc.add_page_break()
    for b in blocks:
        k=b['kind'];text=b.get('text','')
        if k=='table':
            tb=doc.add_table(rows=1,cols=len(b['heads']));tb.style='Table Grid';tb.autofit=False
            widths=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
            for c,w in zip(tb.columns,widths):c.width=Cm(17.2*w)
            for rownum,values in enumerate([b['heads']]+b['rows']):
                row=tb.rows[0] if rownum==0 else tb.add_row()
                cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
                if rownum==0:
                    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                for cell,v,w in zip(row.cells,values,widths):
                    cell.width=Cm(17.2*w);cell.text=v
                    if rownum==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');cell._tc.get_or_add_tcPr().append(sh)
                    for par in cell.paragraphs:
                        par.paragraph_format.space_after=Pt(4);par.paragraph_format.space_before=Pt(3)
                        par.paragraph_format.line_spacing=1.2;fmt_runs(par,9.5,rownum==0)
            continue
        if k=='image':
            par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after=Pt(2);par.add_run().add_picture(str(ASSETS/b['file']),width=Cm(b['width']))
            continue
        style={'title':'Title','h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(k)
        par=doc.add_paragraph(text,style=style)
        if k in ['caption','small','meta']:
            fmt_runs(par,9 if k!='meta' else 10,color='555555' if k!='meta' else '171717')
            par.paragraph_format.line_spacing=1.2;par.paragraph_format.space_after=Pt(4)
        if k=='note':
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F3F4');par._p.get_or_add_pPr().append(sh)
            fmt_runs(par,10);par.paragraph_format.line_spacing=1.25
doc.core_properties.title=STEM;doc.core_properties.author='龙君昱';doc.core_properties.subject='VLA 视角利用研究双周进展'
doc.save(OUT/(STEM+'.docx'))

# Audit text coverage and package integrity; PDF is rendered from the same blocks,
# not an assertion of Word's platform-dependent pagination.
import zipfile
with zipfile.ZipFile(OUT/(STEM+'.docx')) as z:assert z.testzip() is None
import unicodedata
texts=unicodedata.normalize('NFKC','\n'.join(pg.get_text() for pg in pdf))
for required in ['龙君昱','26B951131','64.70','75.88','一、这两周的主要工作','三、后续工作','扰动','熟悉','渲染']:
    assert required in texts, required
assert len(pdf)==3
assert '事务型工作' not in texts
for i,pg in enumerate(pdf):pg.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(ASSETS/f'page-{i+1}.png')
from PIL import Image,ImageOps,ImageDraw
thumbs=[]
for i in range(len(pdf)):
    im=Image.open(ASSETS/f'page-{i+1}.png').convert('RGB');im.thumbnail((358,506));thumbs.append(im)
contact=Image.new('RGB',(3*378,536),'#ddd')
for i,im in enumerate(thumbs):contact.paste(im,((i%3)*378+10,(i//3)*536+10))
contact.save(ASSETS/'contact.png')
manifest={'period':['2026-08-27','2026-09-09'],'report_date':'2026-09-09','revision_date':'2026-09-10',
 'sources':[{'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sources],
 'progress_snapshot':snapshot,'historical_plot_values':dict(zip(keys,vals)),
 'pdf_layout':qa,'docx_checks':['zip integrity','shared content blocks','explicit 3-page sections'],
 'pdf_note':'Same-content typeset reading copy; not a LibreOffice conversion of the DOCX.',
 'new_experiments_started':False,'heldout_outcomes_read':False,
 'outputs':{ext:hashlib.sha256((OUT/(STEM+'.'+ext)).read_bytes()).hexdigest() for ext in ['docx','pdf']}}
(OUT/'_provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'outputs':[str(OUT/(STEM+'.'+x)) for x in ['docx','pdf']], 'layout':qa},ensure_ascii=False,indent=2))
