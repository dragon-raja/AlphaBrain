"""Typeset the 2026-09-10--23 evidence snapshot; no experiment side effects.

PDF, editable DOCX and Markdown share the same content blocks. Run with the
isolated document dependencies, not by installing packages into a GPU runtime.
"""

from pathlib import Path
import argparse
import hashlib
import html
import json
import shutil
import csv


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    receipt = out / '_provenance.json'
    if receipt.exists() and json.loads(receipt.read_text()).get('revision') != 'causes_and_baseline':
        archived = out / 'archive' / 'before_causes_revision'
        if not archived.exists():
            archived.mkdir(parents=True)
            for prior in out.iterdir():
                if prior.is_file():
                    shutil.copy2(prior, archived / prior.name)
            if (out / '_assets').exists():
                shutil.copytree(out / '_assets', archived / '_assets')
    assets = out / '_assets'
    assets.mkdir(exist_ok=True)
    root = Path(__file__).resolve().parents[2]
    sources = [root / 'docs/dsol_paper1' / p for p in [
        'reports/current/vla_view_research_progress_zh.md',
        'results/libero_robocasa_attribution_20260923_zh.md',
        'results/accel_field_20260915_zh.md',
        'history/protocols/local_view_derivative_20260916_zh.md',
        'history/planning/outcome_active_observation_20260918_zh.md',
        'history/planning/active_camera_feasibility_20260921_zh.md',
        'audits/libero_step_zero_interpretation_20260923_zh.md',
    ]]
    texts = {p.name: p.read_text() for p in sources}
    attribution = texts['libero_robocasa_attribution_20260923_zh.md']
    for value in ['65.43', '86.52', '43.50', '37.25', '3168', '0.864647', '0.346834']:
        assert value in attribution, f'Source snapshot missing {value}'

    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties, fontManager
    import numpy as np
    import pymupdf as fitz
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    fontdir = Path('/usr/share/fonts/opentype/noto')
    font = FontProperties(fname=str(fontdir / 'NotoSansCJK-Regular.ttc'))
    fontManager.addfont(str(fontdir / 'NotoSansCJK-Regular.ttc'))
    plt.rcParams.update({'font.family': font.get_name(), 'axes.unicode_minus': False,
                         'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False})
    fig, axes = plt.subplots(1, 2, figsize=(8.6, 2.5), sharey=True)
    chart_values = [np.array([[90.625, 22.406, 65.430], [88.672, 76.459, 86.523]]),
              np.array([[51.25, 40.26, 43.50], [44.88, 40.14, 37.25]])]
    labels = ['规范安装', '扰动后不调整', '最低平均 Accel']
    colors = ['#9BA7B1', '#D89050', '#287D8E']
    for ax, data, title in zip(axes, chart_values, ['LIBERO｜16 个评价初态', 'RoboCasa｜200 个 heldout 初态']):
        x = np.arange(2)
        for j in range(3):
            bars = ax.bar(x + (j - 1) * .24, data[:, j], .22, color=colors[j], label=labels[j], zorder=3)
            for b in bars:
                ax.text(b.get_x() + b.get_width()/2, b.get_height() + 1.6,
                        f'{b.get_height():.1f}', ha='center', fontsize=8.5)
        ax.set_xticks(x, ['规范后训练', '多视角后训练'])
        ax.set_title(title, fontsize=11, pad=10)
        ax.set_ylim(0, 105)
        ax.set_yticks([0, 25, 50, 75, 100])
        ax.grid(axis='y', alpha=.18, zorder=0)
    axes[0].set_ylabel('完整任务成功率（%）')
    fig.legend(*axes[0].get_legend_handles_labels(), loc='lower center', ncol=3,
               frameon=False, bbox_to_anchor=(.51, -.01), fontsize=9)
    fig.subplots_adjust(left=.075, right=.995, top=.84, bottom=.22, wspace=.15)
    fig.savefig(assets / 'view_recovery.png', dpi=200)
    plt.close(fig)

    # Reuse admitted evaluation arrays and the existing scientific plot primitive.
    from AlphaBrain.research.dsol.plotting.robocasa_results import space_axes
    analysis_root = Path('/workspace/research-reviews/robocasa-p0-analysis-20260923')
    population = json.loads((analysis_root / 'population.json').read_text())
    target = np.array([j['split'] == 'target' for j in population['jobs']])
    params = np.array([v['parameters'] for v in population['views']])
    with np.load(analysis_root / 'matrix.npz') as data:
        success = data['success'].mean(-1)[:, target].mean(1)
        accel = data['accel_prefixes'][..., 1].mean(-1)[:, target].mean(1)
    assert target.sum() == 200 and success.shape == accel.shape == (2, 97)
    sources += [analysis_root / 'population.json', analysis_root / 'matrix.npz',
                analysis_root / 'spatial-field-correlation-20260923.json']
    spatial_reference = json.loads(sources[-1].read_text())['benchmarks']['robocasa']['models']
    for m, name in enumerate(('canonical', 'broad')):
        np.testing.assert_allclose(success[m, 1:], spatial_reference[name]['mean_success'])
        np.testing.assert_allclose(accel[m, 1:], spatial_reference[name]['mean_accel'])
    fig = plt.figure(figsize=(8.6, 5.0))
    for m, label in enumerate(['规范后训练', '多视角后训练']):
        upper = space_axes(fig, [.015+.50*m, .535, .395, .375], params[1:], success[m, 1:],
                           label+'｜成功率', percentage=True, canonical=success[m, 0])
        lower = space_axes(fig, [.015+.50*m, .035, .395, .375], params[1:], accel[m, 1:],
                           label+'｜Accel', canonical=accel[m, 0])
        # The caption defines the z coordinate; keep it clear of the colorbar.
        upper.set_zlabel('')
        lower.set_zlabel('')
    fig.savefig(assets / 'accel_space.png', dpi=240)
    plt.close(fig)

    from AlphaBrain.research.dsol.plotting.rig_geometry import compact_candidate_figure
    geometry_root = Path('/workspace/research-reviews/view-geometry-visibility-20260920/robocasa')
    geometry_summary = json.loads((geometry_root / 'summary.json').read_text())
    chosen_initial = geometry_summary['illustration_initial_id']
    with (geometry_root / 'camera_mounts.csv').open() as stream:
        records = [r for r in csv.DictReader(stream) if r['initial_id'] == chosen_initial]
    indexed = {(r['view_id'], r['camera']): r for r in records}
    camera_names = ['robot0_agentview_left', 'robot0_agentview_right']
    ordered = [[indexed[(v['view_id'], c)] for c in camera_names] for v in population['views']]
    positions = np.array([[[float(r[k]) for k in ('x_m','y_m','z_m')] for r in pair] for pair in ordered])
    forwards = np.array([[[float(r[k]) for k in ('forward_x','forward_y','forward_z')] for r in pair] for pair in ordered])
    fig = compact_candidate_figure(params, positions, forwards)
    fig.savefig(assets / 'candidate_geometry.png', dpi=240)
    plt.close(fig)
    sources += [geometry_root / 'camera_mounts.csv', geometry_root / 'summary.json',
                Path('/share/longjunyu/alphabrain/experiments/dsol-robocasa-baseline-panel-20260923/analysis/report.md')]

    # Presentation schematic; these boxes depict the frozen training protocol.
    from matplotlib.patches import FancyBboxPatch
    fig, ax = plt.subplots(figsize=(8.6, 1.8))
    ax.set_axis_off()
    def box(x, y, w, h, label, color):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.012',
                                   facecolor=color,edgecolor='#c2ced6'))
        ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=10.5)
    box(.015,.25,.27,.51,'同一官方 Human300 权重\n同一批 33,000 源窗口\n300 任务；状态／动作标签一致','#edf2f5')
    box(.37,.54,.28,.38,'规范组\n原安装图像 × 2 次曝光','#e8eff8')
    box(.37,.04,.28,.38,'多视角组\n每源示范分配 2 个扰动安装','#e5f3ef')
    box(.735,.25,.25,.51,'各 66,000 样本\n各 8,250 步；全局 batch 8\n训练预算与源数据匹配','#edf2f5')
    for start, end in [((.29,.50),(.355,.73)),((.29,.50),(.355,.23)),
                       ((.665,.73),(.72,.50)),((.665,.23),(.72,.50))]:
        ax.annotate('',xy=end,xytext=start,arrowprops={'arrowstyle':'->','color':'#657c8a','lw':1.5})
    fig.subplots_adjust(left=.005,right=.995,top=.98,bottom=.04)
    fig.savefig(assets / 'training_design.png',dpi=220)
    plt.close(fig)

    external_figures = {
        'training_views.png': Path('/workspace/research-reviews/view-evidence-diagnostic-20260923/contact-sheets/robocasa-OpenDrawer.png'),
        'observation_response.png': Path('/workspace/research-reviews/view-evidence-diagnostic-20260923/observation-response.png'),
    }
    for name, path in external_figures.items():
        shutil.copy2(path, assets / name)
        sources.append(path)
    from AlphaBrain.research.dsol.plotting.observation_evidence import compact_response_figure
    response_source = Path('/workspace/research-reviews/view-evidence-diagnostic-20260923/response-summary.json')
    fig = compact_response_figure(json.loads(response_source.read_text()))
    fig.savefig(assets / 'observation_response_compact.png', dpi=220)
    plt.close(fig)
    sources.append(response_source)

    pages = [[], [], []]
    def add(p, kind, text='', **kwargs):
        pages[p].append(dict(kind=kind, text=text, **kwargs))
    def table(p, heads, rows, widths):
        add(p, 'table', heads=heads, rows=rows, widths=widths)

    add(0, 'title', '科研进展双周报')
    add(0, 'meta', '汇报人：龙君昱　　学号：26B951131')
    add(0, 'meta', '汇报周期：2026 年 9 月 10 日—9 月 23 日　｜　整理日期：9 月 24 日')
    add(0, 'h1', '一、本期完成：RoboCasa 多视角训练与闭环评测')
    add(0, 'image', file='training_design.png')
    add(0, 'caption', '图 1　两组使用相同数据和训练预算，只改变外部相机安装。多视角组每个源窗口用两种视角，不是每帧都用 64 种；腕部图、状态和动作标签不变。')
    add(0, 'image', file='candidate_geometry.png')
    add(0, 'caption', '图 2　97 个候选＝规范 1＋训练 64＋留出 32。左右相机一起调整，位置和朝向都会变；箭头表示朝向。右图为一个初态的安装坐标，非运动轨迹。安装选定后不再调整，但相机会随底盘运动。候选未全部验证真机可达性。')
    add(0, 'p', '两台八卡机完成 400 初态 × 97 视角 × 2 模型 × 4 次执行＝310,400 次完整任务；每个候选另算 8 次 Accel。评价含 10 个已知原子任务，熟悉／留出场景各 200 初态；留出的是布局和风格，不是新任务。')
    add(0, 'h1', '主要现象：RoboCasa 的视角恢复收益较小')
    add(0, 'image', file='view_recovery.png')
    add(0, 'caption', '图 3　先按初态选安装，再执行完整任务。LIBERO 用 16 个评价初态、每候选 32 次；RoboCasa 用 200 个留出初态、每候选 4 次。扰动成绩平均 32 个留出安装；Accel 从全部 97 个候选中选。')
    add(0, 'note', '最低 Accel 的恢复收益：LIBERO 为 +43.02／+10.06 个百分点，RoboCasa 为 +3.24／−2.89（规范／多视角模型）。需要分别解释：原本可恢复的损失有多大，以及 Accel 能否找到有效视角。')

    add(1, 'h1', '二、为什么恢复小？原模型的视角敏感性已经不同')
    table(1, ['后训练前的模型', '规范视角', '四扰动平均', '下降'], [
        ['LIBERO 原生参考', '92.97%', '21.88%', '71.09 pp'],
        ['RoboCasa 官方 Human300', '51.25%', '46.56%', '4.69 pp'],
    ], [.43,.19,.19,.19])
    add(1, 'caption', '独立的小规模基线对照：LIBERO 32 初态；RoboCasa 20 个留出初态；各取 4 个扰动。两边扰动范围不同，不能视为只换 benchmark 的严格对照。')
    add(1, 'p', 'LIBERO 原模型在规范视角下大多会做，扰动后大幅下降；RoboCasa 原模型对所测扰动的平均损失已较小。因此，后训练不充分不能解释全部差别；仅找回规范观察，也无法解决规范视角下原本就有的失败。')
    add(1, 'p', '可能原因是 Human300 多场景训练、底盘运动带来的自然视角变化，以及双外部相机、腕部和本体状态提供了替代信息。各因素尚未分离；自然运动也不等于见过任意安装偏置。平均扰动损失小，不代表所有视角选择的上限小。')
    add(1, 'h1', '三、为什么 Accel 变弱？低值区域也发生了改变')
    add(1, 'image', file='accel_space.png')
    add(1, 'caption', '图 4　RoboCasa 200 个留出初态的平均空间图，96 个扰动点，星号为规范安装。上排越亮成功率越高，下排越亮 Accel 越低；色条各自缩放。坐标同图 2 左图，全部是实测点。Accel 在初态计算，不是整条任务的加速度。这里只比较两份后训练模型，不表示 Accel 低于官方起点。')
    table(1, ['同一候选总体', '规范后训练', '多视角后训练'], [
        ['Accel 均值', '0.2597', '0.2445'],
        ['空间秩相关：−Accel 与成功率', '0.865', '0.347'],
    ], [.54,.23,.23])
    add(1, 'note', '同一 RoboCasa 上，规范模型仍相关，因此不能只用“环境复杂”解释。多视角训练后，差区域有所改善、好区域部分退化；同时 Accel 低值区改变更多：两模型成功率排名相关为 0.753，Accel 排名相关仅 0.407。不是把分数整体压低这么简单。')

    add(2, 'h1', '四、原因分析：训练改变了什么，哪些解释仍需验证')
    add(2, 'p', '当前解释是：多视角训练要求“同一状态换视角后仍预测同一动作”，模型确实变得更不受相机变化影响。但训练没有要求“正确动作对应低 Accel”。因此，去噪变化变小，既可能是更好地适应了视角，也可能是稳定地产生不合适的动作。')
    add(2, 'image', file='observation_response_compact.png')
    add(2, 'caption', '图 5　3,168 次动作生成检查。虚线 1 表示两模型相同。换相机的响应减小，但真实状态响应保留；误差比小于 1 表示改善。RoboCasa 有 7/10 个抽样任务误差下降，整体均值却上升，主要由关冰箱样本拉高。动作误差不等于闭环成功率。')
    add(2, 'p', '已核查：重算 Accel 一致；换用相对排名、只保留输出动作维度、计算完整去噪过程，弱关联仍在。抽样未发现标签错配。换句话说，问题不像是简单的计算错误，但具体训练原因还没有定案。')
    add(2, 'h2', '为什么两边结果不同？以下是有依据、但尚未逐项验证的解释')
    add(2, 'p', '① 模型起点与输入：RoboCasa 已在 Human300 上训练 75k 步，并使用左右外部相机、腕部和本体状态；LIBERO 本轮不输入本体状态。冗余信息可能减弱外部视角影响，但还不能认定存在“腕部／状态捷径”。')
    add(2, 'p', '② 训练覆盖与配置：两边总曝光约 6.5 万，但 LIBERO 分给 8 任务，RoboCasa 分给 300 任务，平均每任务约 8,000 对 220。后者每窗口仅两种视角，且替换规范图；可能覆盖不足或牺牲部分规范表现。两边实现、动作表示和可训练模块也不同，不能只按步数判断训够了。')
    add(2, 'p', '③ 任务与执行过程：RoboCasa 规范成功率本就不高，部分失败可能发生在接触、移动或后续控制；每次执行 25 步再重规划，LIBERO 为 5 步。初态 Accel 更难概括后续变化，但这不能单独解释规范／多视角模型的相关性差异。')
    add(2, 'p', '④ 扰动与可见信息：RoboCasa 方位／仰角为 ±30°／±15°，LIBERO 约 ±60°／±25°，且相机结构不同。前者双相机可互补，但部分把手和接触点仍可能被挡住。增强可能混合了信息充足和不足的样本；标签正确不等于看得清，也不能直接称为脏数据。')
    add(2, 'h1', '五、下一步')
    add(2, 'p', '继续分析低 Accel 失败、较高 Accel 成功的案例，并结合已有中间权重追踪变化。先判断该改训练覆盖、视觉证据利用，还是评分目标，再选一个小实验验证。继续阅读主动感知工作，比较它们如何定义观察收益和训练目标；暂不确定切换框架。')
    add(2, 'h1', '六、本期论文阅读')
    add(2, 'p', 'SaPaVe：相机与操作动作分开建模、两阶段训练。\nActiveArena：主动感知任务集、记忆机制及泛化评价。\nAAWR：用训练期特权信息估计价值，再学习主动感知行为。\nEgoAVFlow：从人类第一视角视频学习，结合 3D flow 和可见性优化视角。')
    add(2, 'small', '文献编号依次为 arXiv:2603.12193、2609.24124、2512.01188、2602.22461。这里只记录阅读内容，尚未用于本期实验。')

    def esc(s):
        return html.escape(str(s)).replace('\n', '<br/>')
    def block_html(b):
        kind, text = b['kind'], esc(b['text'])
        if kind == 'image':
            return f'<img src="{b["file"]}" style="width:100%">'
        if kind == 'table':
            head = ''.join(f'<th style="width:{w*100}%">{esc(t)}</th>' for t,w in zip(b['heads'], b['widths']))
            rows = ''.join('<tr>'+''.join(f'<td style="width:{w*100}%">{esc(v)}</td>' for v,w in zip(row,b['widths']))+'</tr>' for row in b['rows'])
            return '<table><tr>'+head+'</tr>'+rows+'</table>'
        if kind in ('h1', 'h2'):
            return f'<{kind}>{text}</{kind}>'
        return f'<p class="{kind}">{text}</p>'

    css = '''
    @font-face {font-family:cn;src:url(NotoSansCJK-Regular.ttc);}
    @font-face {font-family:cn;src:url(NotoSansCJK-Bold.ttc);font-weight:bold;}
    body {font-family:cn;font-size:10pt;line-height:1.42;color:#1d252c;}
    p{margin:0 0 6pt;} .title{font-size:20pt;font-weight:bold;margin-bottom:12pt;}
    .meta{font-size:9.2pt;color:#46515b;margin-bottom:3pt;}
    h1{font-size:13.5pt;margin:12pt 0 7pt;color:#183b51;}
    h2{font-size:10.5pt;margin:9pt 0 5pt;font-weight:bold;}
    table{width:100%;border-collapse:collapse;table-layout:fixed;margin:3pt 0 8pt;}
    th,td{border:.5pt solid #cbd3d9;padding:5pt 6pt;vertical-align:top;font-size:9.2pt;line-height:1.35;}
    th{background:#edf2f5;text-align:left;font-weight:bold;}
    .caption{font-size:8.1pt;line-height:1.35;color:#59636b;margin:3pt 0 7pt;}
    .note{font-size:9.6pt;background:#edf3f6;padding:7pt;line-height:1.4;margin-top:8pt;}
    .small{font-size:7.6pt;color:#667078;line-height:1.35;margin-top:7pt;}
    '''
    archive = fitz.Archive([str(assets), str(fontdir)])
    stem = '科研进展双周报_20260910-20260923'
    pdf = fitz.open()
    layout = []
    for i, blocks in enumerate(pages):
        pg = pdf.new_page(width=595.28, height=841.89)
        spare, scale = pg.insert_htmlbox(fitz.Rect(51, 41, 544, 792), ''.join(map(block_html, blocks)),
                                        css=css, archive=archive, scale_low=1)
        if spare < 0:
            raise ValueError(f'Page {i+1} overflow; edit content instead of shrinking')
        pg.insert_htmlbox(fitz.Rect(51, 807, 530, 830),
                          f'科研进展双周报 · VLA 视角利用 · 2026.09.10—09.23　　　　　　　　　{i+1} / {len(pages)}',
                          css=css+'body{font-size:7.5pt;color:#737d86;}', archive=archive)
        layout.append(dict(page=i+1, spare_height=spare, scale=scale))
    pdf.set_metadata({'title':stem, 'author':'龙君昱', 'subject':'VLA视角利用：跨环境验证、问题分析与后续方向'})
    pdf.save(out / (stem+'.pdf'), garbage=4, deflate=True)
    extracted = '\n'.join(p.get_text() for p in pdf)
    for s in ['310,400', '3,168', '92.97', 'ActiveArena', '五、下一步']:
        assert s in extracted, s
    for i, pg in enumerate(pdf):
        pg.get_pixmap(matrix=fitz.Matrix(1.4,1.4)).save(assets / f'page-{i+1}.png')

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.45), Cm(1.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    for name, size in [('Normal',10), ('Title',20), ('Heading 1',13.5), ('Heading 2',10.5)]:
        st = doc.styles[name]
        st.font.name, st.font.size = 'Arial', Pt(size)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '等线')
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_after = Pt(5)
        if name != 'Normal':
            st.font.color.rgb = RGBColor.from_string('183B51')
            st.paragraph_format.space_before = Pt(8)
    for pi, blocks in enumerate(pages):
        if pi:
            doc.add_page_break()
        for b in blocks:
            k = b['kind']
            if k == 'table':
                tb = doc.add_table(rows=1, cols=len(b['heads']))
                tb.style = 'Table Grid'
                tb.autofit = False
                for col, w in zip(tb.columns, b['widths']):
                    col.width = Cm(17.4*w)
                for ri, values in enumerate([b['heads']]+b['rows']):
                    row = tb.rows[0] if ri == 0 else tb.add_row()
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell,v,w in zip(row.cells,values,b['widths']):
                        cell.width = Cm(17.4*w)
                        cell.text = v
                        if ri == 0:
                            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF2F5');cell._tc.get_or_add_tcPr().append(sh)
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after=Pt(4)
                            p.paragraph_format.line_spacing=1.15
                            for run in p.runs:
                                run.font.size=Pt(9)
                                run.bold=ri==0
                continue
            if k == 'image':
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(assets/b['file']), width=Cm(17.2))
                continue
            p=doc.add_paragraph(b['text'], style={'title':'Title','h1':'Heading 1','h2':'Heading 2'}.get(k))
            if k in ['meta','caption','small']:
                for run in p.runs:
                    run.font.size=Pt({'meta':9,'caption':8,'small':7.5}[k])
                    run.font.color.rgb=RGBColor.from_string('59636B')
            if k == 'note':
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF3F6');p._p.get_or_add_pPr().append(sh)
    sec.footer.paragraphs[0].text='科研进展双周报 · VLA 视角利用 · 2026.09.10—09.23'
    doc.core_properties.title=stem
    doc.core_properties.author='龙君昱'
    doc.save(out/(stem+'.docx'))
    md=[]
    for blocks in pages:
        for b in blocks:
            if b['kind']=='table':
                md += ['| '+' | '.join(b['heads'])+' |','|'+'---|'*len(b['heads'])]
                md += ['| '+' | '.join(row)+' |' for row in b['rows']]
            elif b['kind']=='image':
                md.append(f'![研究图表：{b["file"]}](_assets/{b["file"]})')
            else:
                prefix={'title':'# ','h1':'## ','h2':'### '}.get(b['kind'],'')
                md.append(prefix+b['text'])
            md.append('')
    (out/(stem+'.md')).write_text('\n'.join(md))
    record = dict(revision='causes_and_baseline', period=['2026-09-10','2026-09-23'], compiled='2026-09-24',
                  thread='01a0861a-5f0e-79b1-9cf9-b1e1ede4baf2',
                  sources=[dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in sources],
                  chart_data=[v.tolist() for v in chart_values], layout=layout,
                  references=['https://arxiv.org/abs/2603.12193', 'https://arxiv.org/abs/2609.24124',
                              'https://arxiv.org/abs/2512.01188', 'https://arxiv.org/abs/2602.22461'],
                  notes=['No new experiments; no current research report overwritten.',
                         'DOCX shares content; Word pagination may depend on installed fonts.',
                         'ActiveArena resource review remains prospective; no local reproduction claimed.'])
    record['outputs']={suffix:hashlib.sha256((out/(stem+suffix)).read_bytes()).hexdigest() for suffix in ['.pdf','.docx','.md']}
    (out/'_provenance.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(output=str(out),layout=layout),ensure_ascii=False))


if __name__ == '__main__':
    main()
