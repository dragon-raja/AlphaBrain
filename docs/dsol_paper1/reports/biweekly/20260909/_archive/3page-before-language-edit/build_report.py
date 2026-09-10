"""Read-only research evidence -> editable biweekly report and same-content PDF.

No policy execution, training, or evaluation-controller modification.
Dependencies are isolated in /tmp, not installed into the experiment environment.
"""
from pathlib import Path
import sys, json, hashlib, html
sys.path[:0] = ['/tmp/dsol-weekly-docx-deps-20260909',
                '/root/.cache/uv/archive-v0/XO8_y4zBEt33DOY7wfN64']
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, fontManager
import pymupdf as fitz
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent
ASSETS = OUT / '_assets'
ASSETS.mkdir(exist_ok=True)
REPO = Path('/workspace/projects/alphabrain-dsol-paper1')
EXP = Path('/share/longjunyu/alphabrain/experiments')
NEW = EXP / 'dsol-standard-initial-aa0-v1-20260908'
OLD = EXP / 'dsol-view-landscape-v1-20260908'
STEM = '科研进展双周报_20260827-20260909'
font = FontProperties(fname='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
plt.rcParams.update({'font.family': font.get_name(), 'axes.unicode_minus': False,
                     'font.size': 11, 'axes.spines.top': False, 'axes.spines.right': False})

sources = [
    REPO/'docs/dsol_paper1/training_match_execution_status_20260907_zh.md',
    REPO/'docs/dsol_paper1/libero_expanded_a_constructed_gate_20260827_zh.md',
    REPO/'docs/dsol_paper1/view_repeatability_root_cause_20260908_zh.md',
    REPO/'docs/dsol_paper1/render_protocol_bridge_and_historical_validity_20260908_zh.md',
    REPO/'docs/dsol_paper1/standard_initial_aa0_execution_20260908_zh.md',
    OLD/'checkpoint-audit/completed_run_receipt.json',
    OLD/'metric-oracle-comparison-v2/broad_summary.json',
    OLD/'render-protocol-bridge-v1/analysis-recovered-v1.json',
    NEW/'release.json', NEW/'assets/aa0-a/00/policy_inputs.npz',
]

rules = json.loads(sources[6].read_text())['overall']['rules']
keys = ['canonical', 'min_accel', 'accel_top10_max_visibility', 'statewise_empirical_max_O32']
vals = [rules[k]['success']['mean']*100 for k in keys]
assert np.allclose(vals, [64.697265625,61.376953125,60.498046875,75.87890625])

# Historical spatial distributions are auxiliary illustrations, not new outcomes.
sources += [OLD/'broad-existing/verified_arrays.npz',
            OLD/'broad-existing/candidate_geometry.json']
with np.load(sources[-2]) as a:
    success = a['success'].mean((0,2))
    accel = a['accel'].mean((0,2))
    visibility = a['visibility'].mean(0)
    candidate_ids = a['candidate_ids'].tolist()
poses_by_id = {x['pose_id']:x for x in json.loads(sources[-1].read_text())['poses']}
poses = [poses_by_id[x] for x in candidate_ids]
xyz=np.array([[x[k] for k in ['azimuth_deg','elevation_deg','radius_scale']] for x in poses])
fig=plt.figure(figsize=(8,3.1))
for j,(v,title,cmap) in enumerate([
    (success*100,'续跑成功率（%）','viridis'),
    (accel,'平均 Accel','magma'),
    (visibility*100,'任务实体可见像素占比（%）','cividis')]):
    ax=fig.add_subplot(1,3,j+1,projection='3d')
    low,high=(0,100) if j==0 else (float(v.min()),float(v.max()))
    sc=ax.scatter(*xyz[1:].T,c=v[1:],s=16,cmap=cmap,vmin=low,vmax=high,depthshade=False)
    ax.scatter(*xyz[0],c=[v[0]],s=65,marker='*',edgecolor='black',
               cmap=cmap,vmin=low,vmax=high,depthshade=False)
    ax.set_title(title,fontsize=9,pad=4)
    ax.set_xlabel('方位角偏移',fontsize=7,labelpad=-3)
    ax.set_ylabel('仰角偏移',fontsize=7,labelpad=-3)
    ax.set_zlabel('半径倍率',fontsize=7,labelpad=-7)
    ax.tick_params(labelsize=6,pad=-1);ax.view_init(22,-58)
    cb=fig.colorbar(sc,ax=ax,orientation='horizontal',shrink=.72,pad=.13,fraction=.04)
    cb.ax.tick_params(labelsize=6)
fig.subplots_adjust(left=.01,right=.98,bottom=.13,top=.92,wspace=.13)
fig.savefig(ASSETS/'historical_spatial_maps.png',dpi=220);plt.close(fig)

# Freeze a progress snapshot; intentionally do not read held-out outcome labels.
snapshot_file = OUT/'_progress_snapshot.json'
if snapshot_file.exists():
    snapshot=json.loads(snapshot_file.read_text())
else:
    from datetime import datetime, timezone
    snapshot={'observed_at_utc':datetime.now(timezone.utc).isoformat(),'hosts':{}}
    for host in ['fresh','gnu']:
        d=NEW/'hosts'/host
        snapshot['hosts'][host]={
            'status':json.loads((d/'status.json').read_text()),
            'first_wave':json.loads((d/'dense-init-0-complete.json').read_text()),
            'initial_gate':json.loads((d/'gate-pass.json').read_text()),
            'dynamic_gate':json.loads((NEW/'scheduling/dynamic-v1'/host/'cross-host-gate-pass.json').read_text()),
        }
    snapshot_file.write_text(json.dumps(snapshot,ensure_ascii=False,indent=2)+'\n')
assert all(v['first_wave']['episodes']==24832 for v in snapshot['hosts'].values())
from datetime import datetime, timedelta
asof = (datetime.fromisoformat(snapshot['observed_at_utc'])+timedelta(hours=8)).strftime('%m 月 %d 日 %H:%M')

pages=[]
def page():
    x=[];pages.append(x);return x
def add(p,kind,text='',**kw):p.append(dict(kind=kind,text=text,**kw))
def table(p,heads,rows,widths=None):add(p,'table',heads=heads,rows=rows,widths=widths)


p=page()
add(p,'title','科研进展周报')
add(p,'meta','汇报人：____________    学号：____________')
add(p,'meta','统计周期：2026 年 8 月 27 日—9 月 9 日（双周）')
add(p,'meta',f'研究方向：VLA 的视角利用    进度截点：北京时间 {asof}')
add(p,'h1','一、研究重点与阶段进展')
add(p,'p','当前研究重点是：相机位姿发生扰动后，能否通过模型推理指标及视觉特征，在候选空间中识别模型较适应的视角，恢复任务执行能力；并比较规范视角训练与多视角训练对这一过程的影响。研究不以找到 Oracle 最优视角为唯一目标，而是先分析“什么样的视角更容易被模型有效利用”。')
table(p,['待验证问题','对应研究思路'],[
 ['熟悉视角能否帮助恢复性能？','对照候选到规范位姿、训练视角集合的距离，以及 Accel 等指标，检验从扰动视角选回模型较适应区域后，整任务成功率能否改善。'],
 ['训练覆盖之后，什么信息仍有价值？','比较两种训练下的成功率空间分布；在几何距离相近的候选中，分析可见性差异是否对应额外收益，而非仅返回训练过的位置。'],
],[.30,.70])
add(p,'note','“熟悉度”目前是待检验的解释，不是已测得的模型能力：到训练位姿的距离只是几何代理，低 Accel 也不能直接等同于熟悉。后续需用闭环成功率检验，并区分几何接近与图像内容相似。')
add(p,'h2','1. 完成两种训练条件的匹配对照')
add(p,'p','以已有 Broad64 模型为锚，完成规范视角模型补训：同一 π0.5 初始化、数据记录和划分，保留相同腕部图、语言及动作监督，只改变外部训练图像；两侧均为 seed 41、2,000 次优化更新、64,000 次样本呈现，学习率轨迹匹配。该对照用于分析训练覆盖如何改变视角敏感性，当前不推广为多训练 seed 结论。')
add(p,'h2','2. 统一评测协议，支撑指标与成功率直接对应')
add(p,'p','重复性检查将最早分歧定位到 MuJoCo/EGL 多重采样抗锯齿渲染。完成 1,536 次新旧协议配对复核后，主评测统一采用关闭多重采样的 AA0，并检查候选指标图与实际初始策略输入一致。双机重复轨迹及调度优化验收已通过；历史结果保留为辅助证据，不与新协议标签混用。')
add(p,'h2','3. 标准初态整任务矩阵正在执行')
table(p,['评测范围','进度与用途'],[
 ['8 任务 × 4 标准初态 × 97 候选 × 32 套噪声 × 2 模型','总计 198,656 次完整任务执行。第一轮 49,664 次已完成；截至上述截点，两机均在执行第二轮。'],
 ['初态、候选和噪声跨模型配对','每次执行固定一个外部视角；不增加背景变化或构造遮挡。共同矩阵用于空间分布、指标规则和选择器分析。'],
],[.48,.52])

p=page()
add(p,'h1','二、已有结果与指标分析')
add(p,'h2','1. 从单个最好视角转向空间分布分析')
add(p,'p','已对旧 Broad64 数据整理成功率、Accel 与可见性的候选空间分布。接下来重点比较：高成功率区域是否接近训练覆盖，低 Accel 区域是否与之重合，以及多视角训练是否扩大了可成功执行的区域。图 1 保留实际离散坐标，不将插值结果视为实测。')
add(p,'image',file='historical_spatial_maps.png',width=17.0)
add(p,'caption','图 1  历史 Broad64 模型的三类空间分布，颜色为跨 64 个恢复起点的均值。坐标是候选相对规范相机的方位角、仰角及半径参数，星号为规范视角，并非不同场景共用的世界坐标。旧渲染协议结果仅作分析示例；总体均值可能掩盖任务差异，新矩阵将逐任务、逐初态比较。')
add(p,'h2','2. 指标存在利用价值的线索，但尚不能代表可靠视角价值')
table(p,['旧矩阵中的规则','恢复后续跑成功率'],[
 ['97 个候选均匀选择（按矩阵计算的期望）','58.31%'],
 ['最低平均 Accel','61.38%'],
 ['Accel 前 10 名中按可见性选择','60.50%'],
 ['规范视角','64.70%'],
],[.72,.28])
add(p,'p','最低 Accel 的点估计高于均匀选择，但未超过规范视角；加入可见性也未在该总体上进一步改善。这更适合提出“能否避开不适应区域、找回可用视角”的问题，而不是直接认定指标能找到最优候选。尚不能据此宣称相对扰动视角的稳定恢复收益已经成立。')
add(p,'p','可见性构造实验也未形成普遍正收益：旧两任务、12 个恢复状态中，信息视角为 75.0%，规范视角为 83.3%；匹配视角支持补训后，信息视角仍为 75.0%。因此，任务实体像素占比只能作为待验证特征，不能等同于任务所需信息量。')
add(p,'note','本页数字均来自历史恢复状态续跑，不是当前标准初态整任务成绩。候选经验 Oracle 仅用于估计可选收益空间、衡量方法差距；不将超过规范视角或逼近 Oracle 设为“视角利用有价值”的唯一判据。')

p=page()
add(p,'h1','三、下一阶段分析与验证安排')
add(p,'p','后续按“先看分布—再检验指标—最后决定是否学习选择器”的顺序推进，分别检验性能恢复与信息增益，不预设必须采用某一种选择算法。')
table(p,['顺序','具体工作','关键判断'],[
 ['1. 空间分布','在相同初态下并列绘制两模型的成功率、平均 Accel、Accel 跨初噪波动及可见性；按半径分层或三维散点展示，并绘制两模型成功率差值图。','训练改变了哪些区域？指标低值是否对应高成功率？高低分区域是否跨初态稳定？'],
 ['2. 熟悉度代理','分别计算到规范位姿、最近后训练视角的距离；必要时补算冻结视觉特征到训练图像的相似度。按任务及几何距离分层，检查其与指标、成功率的关系。','低 Accel 是否反映训练分布适应性，还是仅与目标大小、遮挡等因素相关？几何接近不直接视为因果证明。'],
 ['3. 扰动后恢复','事前固定扰动起始视角集合；比较保持扰动视角、返回规范视角、最近训练位姿、均匀选择、最低 Accel 与 Accel Top-10 + 可见性。','方法相对扰动起点能恢复多少成功率？是否只是回到规范位置？多视角训练后收益是否缩小？'],
 ['4. 条件性信息收益','在几何熟悉度相近的候选内，比较任务实体可见性、遮挡及目标尺度变化；区分成功率饱和、能力不足和具有辨识力的任务。','在已有训练覆盖下，是否仍有信息变化对应的额外收益？先做关联分析，不将其直接写成因果机制。'],
 ['5. 选择方法','只有指标显示可复现规律后，再考虑组合打分或轻量回归／排序；保留规范候选及回退规则，固定方法后在留出初态验证。','是否优于简单摆放或单指标基线？经验 Oracle 作为辅助参照，不替代实际方法成功率。'],
],[.15,.48,.37])
add(p,'h2','数据复用与验证边界')
add(p,'p','空间图、指标对照和冻结规则的成绩可复用同一完整闭环矩阵；新增视觉特征通常只需补算表示，不必重跑全候选。初态 0、1 用于开发，2、3 用于留出验证，测试成败不参与调参；必要的独立新噪声复评仅面向少数最终方法。')
add(p,'p','当前矩阵检验的是“任务开始前选定视角”的效果，不包含真实相机从扰动位姿移动到目标位姿的时间与观察代价。读取全部候选图计算 Accel 属于候选图可查询条件；如要验证在线搜索和实际移动，需另设小规模获取预算与闭环实验，不能从现有矩阵直接声称已完成主动感知。')
add(p,'h2','预期形成的研究结论')
add(p,'p','围绕两点收敛：一是指标能否帮助模型在视角扰动后找回可用观察、恢复任务能力；二是多视角训练如何改变这种需求，以及何种任务信息变化仍值得利用。当前尚未确立最终算法贡献，先由匹配实验与空间分析决定方法设计，再规划追加初态、RoboCasa 或真机验证。')
add(p,'small','数据截至 2026 年 9 月 9 日；本次仅修订周报结构与研究重点，不调整运行评测。依据：匹配训练审计、历史视角矩阵与可见性实验、渲染配对复核、标准初态冻结协议及双机进度记录。')

def esc(x):return html.escape(str(x)).replace('\n','<br/>')
def html_block(b):
    k=b['kind'];t=esc(b.get('text',''))
    if k=='table':
        ws=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
        head='<tr>'+''.join(f'<th style="width:{w*100:.1f}%">{esc(v)}</th>' for v,w in zip(b['heads'],ws))+'</tr>'
        rows=''.join('<tr>'+''.join(f'<td>{esc(v)}</td>' for v in row)+'</tr>' for row in b['rows'])
        return '<table>'+head+rows+'</table>'
    if k=='image':return f'<div class="image"><img src="{b["file"]}" width="100%"/></div>'
    if k in ['h1','h2','h3']:return f'<{k}>{t}</{k}>'
    return f'<p class="{k}">{t}</p>'

css='''
@font-face {font-family:report;src:url(NotoSansCJK-Regular.ttc);}
@font-face {font-family:report;src:url(NotoSansCJK-Bold.ttc);font-weight:bold;}
body {font-family:report;font-size:10.5pt;line-height:1.48;color:#171717;}
p {margin:0 0 7pt 0;} .title{font-size:20pt;font-weight:bold;margin-bottom:13pt;}
.meta {font-size:10pt;margin-bottom:3pt;} h1{font-size:15pt;margin:13pt 0 8pt;}
h2{font-size:11.5pt;margin:10pt 0 6pt;} h3{font-size:11.5pt;margin:9pt 0 7pt;}
table{width:100%;border-collapse:collapse;margin:5pt 0 9pt;table-layout:fixed;}
th,td{border:0.55pt solid #cdd3d7;padding:5pt 6pt;vertical-align:top;font-size:9.6pt;line-height:1.4;}
th{background:#edf0f2;text-align:left;font-weight:bold;}
.caption{font-size:9pt;color:#555;line-height:1.4;margin:2pt 0 9pt;}
.note{font-size:10pt;background:#f2f3f4;padding:8pt;line-height:1.45;margin-top:7pt;}
.small{font-size:8.5pt;color:#666;margin-top:9pt;}.image{margin:4pt 0;}
'''
pdf=fitz.open();archive=fitz.Archive([str(ASSETS),'/usr/share/fonts/opentype/noto'])
qa=[]
for i,blocks in enumerate(pages):
    pg=pdf.new_page(width=595.28,height=841.89)
    body=''.join(html_block(b) for b in blocks)
    spare,scale=pg.insert_htmlbox(fitz.Rect(54,46,541.28,791),body,css=css,archive=archive,scale_low=1)
    if spare < 0:raise RuntimeError(f'Page {i+1} overflow; shorten content, do not auto shrink')
    pg.insert_htmlbox(fitz.Rect(54,806,470,830),'科研进展周报 · VLA 视角利用 · 2026.08.27—09.09',css=css+'body{font-size:8pt;color:#777;}',archive=archive)
    pg.insert_text((527,818),str(i+1),fontsize=8,color=(.45,.45,.45))
    qa.append({'page':i+1,'spare_height':spare,'scale':scale})
pdf.set_metadata({'title':STEM,'author':'','subject':'双周科研进展；历史辅助证据与当前主评测分开呈现'})
pdf.save(OUT/(STEM+'.pdf'),garbage=4,deflate=True)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Heading 1',15),('Heading 2',11.5),('Heading 3',11.5)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True
    st.font.color.rgb=RGBColor.from_string('171717');st.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0);st.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0]
footer.text='科研进展周报 · VLA 视角利用\t'
footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.8),WD_ALIGN_PARAGRAPH.RIGHT)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('777777')
def fmt_runs(par,size=10.5,bold=False,color='171717'):
    for r in par.runs:
        r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')

for pi,blocks in enumerate(pages):
    if pi:doc.add_page_break()
    for b in blocks:
        k=b['kind'];text=b.get('text','')
        if k=='table':
            tb=doc.add_table(rows=1,cols=len(b['heads']));tb.style='Table Grid';tb.autofit=False
            widths=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
            for c,w in zip(tb.columns,widths):c.width=Cm(17.2*w)
            for rownum,values in enumerate([b['heads']]+b['rows']):
                row=tb.rows[0] if rownum==0 else tb.add_row()
                cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
                if rownum==0:
                    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                for cell,v,w in zip(row.cells,values,widths):
                    cell.width=Cm(17.2*w);cell.text=v
                    if rownum==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');cell._tc.get_or_add_tcPr().append(sh)
                    for par in cell.paragraphs:
                        par.paragraph_format.space_after=Pt(4);par.paragraph_format.space_before=Pt(3)
                        par.paragraph_format.line_spacing=1.2;fmt_runs(par,9.5,rownum==0)
            continue
        if k=='image':
            par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after=Pt(2);par.add_run().add_picture(str(ASSETS/b['file']),width=Cm(b['width']))
            continue
        style={'title':'Title','h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(k)
        par=doc.add_paragraph(text,style=style)
        if k in ['caption','small','meta']:
            fmt_runs(par,9 if k!='meta' else 10,color='555555' if k!='meta' else '171717')
            par.paragraph_format.line_spacing=1.2;par.paragraph_format.space_after=Pt(4)
        if k=='note':
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F3F4');par._p.get_or_add_pPr().append(sh)
            fmt_runs(par,10);par.paragraph_format.line_spacing=1.25
doc.core_properties.title=STEM;doc.core_properties.author='';doc.core_properties.subject='VLA 视角利用研究双周进展'
doc.save(OUT/(STEM+'.docx'))

# Audit text coverage and package integrity; PDF is rendered from the same blocks,
# not an assertion of Word's platform-dependent pagination.
import zipfile
with zipfile.ZipFile(OUT/(STEM+'.docx')) as z:assert z.testzip() is None
texts='\n'.join(pg.get_text() for pg in pdf)
for required in ['198,656','49,664','64.70','一、研究重点','三、下一阶段','扰动','熟悉','渲染']:
    assert required in texts, required
assert len(pdf)==3
assert '事务型工作' not in texts
for i,pg in enumerate(pdf):pg.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(ASSETS/f'page-{i+1}.png')
from PIL import Image,ImageOps,ImageDraw
thumbs=[]
for i in range(len(pdf)):
    im=Image.open(ASSETS/f'page-{i+1}.png').convert('RGB');im.thumbnail((358,506));thumbs.append(im)
contact=Image.new('RGB',(3*378,536),'#ddd')
for i,im in enumerate(thumbs):contact.paste(im,((i%3)*378+10,(i//3)*536+10))
contact.save(ASSETS/'contact.png')
manifest={'period':['2026-08-27','2026-09-09'],'report_date':'2026-09-09','revision_date':'2026-09-10',
 'sources':[{'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sources],
 'progress_snapshot':snapshot,'historical_plot_values':dict(zip(keys,vals)),
 'pdf_layout':qa,'docx_checks':['zip integrity','shared content blocks','explicit 3-page sections'],
 'pdf_note':'Same-content typeset reading copy; not a LibreOffice conversion of the DOCX.',
 'new_experiments_started':False,'heldout_outcomes_read':False,
 'outputs':{ext:hashlib.sha256((OUT/(STEM+'.'+ext)).read_bytes()).hexdigest() for ext in ['docx','pdf']}}
(OUT/'_provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'outputs':[str(OUT/(STEM+'.'+x)) for x in ['docx','pdf']], 'layout':qa},ensure_ascii=False,indent=2))
