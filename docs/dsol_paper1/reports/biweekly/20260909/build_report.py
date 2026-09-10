"""Read-only research evidence -> editable biweekly report and same-content PDF.

No policy execution, training, or evaluation-controller modification.
Dependencies are isolated in /tmp, not installed into the experiment environment.
"""
from pathlib import Path
import sys, json, hashlib, html
sys.path[:0] = ['/tmp/dsol-weekly-docx-deps-20260909',
                '/root/.cache/uv/archive-v0/XO8_y4zBEt33DOY7wfN64']
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, fontManager
import pymupdf as fitz
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent
ASSETS = OUT / '_assets'
ASSETS.mkdir(exist_ok=True)
REPO = Path('/workspace/projects/alphabrain-dsol-paper1')
EXP = Path('/share/longjunyu/alphabrain/experiments')
NEW = EXP / 'dsol-standard-initial-aa0-v1-20260908'
OLD = EXP / 'dsol-view-landscape-v1-20260908'
STEM = '科研进展双周报_20260827-20260909'
font = FontProperties(fname='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
plt.rcParams.update({'font.family': font.get_name(), 'axes.unicode_minus': False,
                     'font.size': 11, 'axes.spines.top': False, 'axes.spines.right': False})

sources = [
    REPO/'docs/dsol_paper1/training_match_execution_status_20260907_zh.md',
    REPO/'docs/dsol_paper1/libero_expanded_a_constructed_gate_20260827_zh.md',
    REPO/'docs/dsol_paper1/view_repeatability_root_cause_20260908_zh.md',
    REPO/'docs/dsol_paper1/render_protocol_bridge_and_historical_validity_20260908_zh.md',
    REPO/'docs/dsol_paper1/standard_initial_aa0_execution_20260908_zh.md',
    OLD/'checkpoint-audit/completed_run_receipt.json',
    OLD/'metric-oracle-comparison-v2/broad_summary.json',
    OLD/'render-protocol-bridge-v1/analysis-recovered-v1.json',
    NEW/'release.json', NEW/'assets/aa0-a/00/policy_inputs.npz',
]

rules = json.loads(sources[6].read_text())['overall']['rules']
keys = ['canonical', 'min_accel', 'accel_top10_max_visibility', 'statewise_empirical_max_O32']
vals = [rules[k]['success']['mean']*100 for k in keys]
assert np.allclose(vals, [64.697265625,61.376953125,60.498046875,75.87890625])

# Historical spatial distributions are auxiliary illustrations, not new outcomes.
sources += [OLD/'broad-existing/verified_arrays.npz',
            OLD/'broad-existing/candidate_geometry.json']
with np.load(sources[-2]) as a:
    success = a['success'].mean((0,2))
    accel = a['accel'].mean((0,2))
    visibility = a['visibility'].mean(0)
    candidate_ids = a['candidate_ids'].tolist()
poses_by_id = {x['pose_id']:x for x in json.loads(sources[-1].read_text())['poses']}
poses = [poses_by_id[x] for x in candidate_ids]
xyz=np.array([[x[k] for k in ['azimuth_deg','elevation_deg','radius_scale']] for x in poses])
fig=plt.figure(figsize=(8,3.1))
for j,(v,title,cmap) in enumerate([
    (success*100,'续跑成功率（%）','viridis'),
    (accel,'平均 Accel','magma'),
    (visibility*100,'任务实体可见像素占比（%）','cividis')]):
    ax=fig.add_subplot(1,3,j+1,projection='3d')
    low,high=(0,100) if j==0 else (float(v.min()),float(v.max()))
    sc=ax.scatter(*xyz[1:].T,c=v[1:],s=16,cmap=cmap,vmin=low,vmax=high,depthshade=False)
    ax.scatter(*xyz[0],c=[v[0]],s=65,marker='*',edgecolor='black',
               cmap=cmap,vmin=low,vmax=high,depthshade=False)
    ax.set_title(title,fontsize=9,pad=4)
    ax.set_xlabel('方位角偏移',fontsize=7,labelpad=-3)
    ax.set_ylabel('仰角偏移',fontsize=7,labelpad=-3)
    ax.set_zlabel('半径倍率',fontsize=7,labelpad=-7)
    ax.tick_params(labelsize=6,pad=-1);ax.view_init(22,-58)
    cb=fig.colorbar(sc,ax=ax,orientation='horizontal',shrink=.72,pad=.13,fraction=.04)
    cb.ax.tick_params(labelsize=6)
fig.subplots_adjust(left=.01,right=.98,bottom=.13,top=.92,wspace=.13)
fig.savefig(ASSETS/'historical_spatial_maps.png',dpi=220);plt.close(fig)


# Experimental comparison with actual candidate observations (not generated images).
from matplotlib.patches import FancyBboxPatch
fig=plt.figure(figsize=(7.4,3.75),facecolor='white')
with np.load(NEW/'assets/aa0-a/00/policy_inputs.npz') as a:
    ids=a['candidate_ids'].tolist()
    for j,(cid,label) in enumerate([
        ('canonical','规范视角'),('broad_train_000','训练目录候选'),('broad_heldout_000','留出目录候选')]):
        ax=fig.add_axes([.055+j*.32,.46,.25,.47])
        ax.imshow(a['external_images'][ids.index(cid)]);ax.axis('off');ax.set_title(label,fontsize=11,pad=4)
canvas=fig.add_axes([0,0,1,1]);canvas.set_axis_off()
canvas.text(.5,.425,'同一任务初态；各候选分别输入，不同时融合',ha='center',fontsize=10,color='#4c5963')
for x,txt in [(.09,'规范视角训练模型'),(.56,'多视角训练模型')]:
    canvas.add_patch(FancyBboxPatch((x,.205),.35,.12,boxstyle='round,pad=.01',facecolor='#edf2f6',edgecolor='#90a6b8'))
    canvas.text(x+.175,.266,txt,ha='center',va='center',fontsize=12,color='#233e52')
    canvas.annotate('',xy=(x+.175,.34),xytext=(.5,.395),arrowprops=dict(arrowstyle='->',color='#6b7a83',lw=1.3))
canvas.text(.5,.11,'比较：整任务成功率  /  Accel  /  任务实体可见性',ha='center',fontsize=11)
canvas.text(.5,.025,'每次执行仅固定一个外部视角，腕部输入保留',ha='center',fontsize=9,color='#666')
fig.savefig(ASSETS/'experiment_overview.png',dpi=190,bbox_inches='tight',pad_inches=.02);plt.close(fig)

# Freeze a progress snapshot; intentionally do not read held-out outcome labels.
snapshot_file = OUT/'_progress_snapshot.json'
if snapshot_file.exists():
    snapshot=json.loads(snapshot_file.read_text())
else:
    from datetime import datetime, timezone
    snapshot={'observed_at_utc':datetime.now(timezone.utc).isoformat(),'hosts':{}}
    for host in ['fresh','gnu']:
        d=NEW/'hosts'/host
        snapshot['hosts'][host]={
            'status':json.loads((d/'status.json').read_text()),
            'first_wave':json.loads((d/'dense-init-0-complete.json').read_text()),
            'initial_gate':json.loads((d/'gate-pass.json').read_text()),
            'dynamic_gate':json.loads((NEW/'scheduling/dynamic-v1'/host/'cross-host-gate-pass.json').read_text()),
        }
    snapshot_file.write_text(json.dumps(snapshot,ensure_ascii=False,indent=2)+'\n')
assert all(v['first_wave']['episodes']==24832 for v in snapshot['hosts'].values())
from datetime import datetime, timedelta
asof = (datetime.fromisoformat(snapshot['observed_at_utc'])+timedelta(hours=8)).strftime('%m 月 %d 日 %H:%M')

pages=[]
def page():
    x=[];pages.append(x);return x
def add(p,kind,text='',**kw):p.append(dict(kind=kind,text=text,**kw))
def table(p,heads,rows,widths=None):add(p,'table',heads=heads,rows=rows,widths=widths)




p=page()
add(p,'title','科研进展周报')
add(p,'meta','汇报人：龙君昱    学号：26B951131')
add(p,'meta','汇报周期：2026 年 8 月 27 日—9 月 9 日')
add(p,'h1','一、主要工作')
add(p,'p','这两周主要研究两件事：相机扰动后，能否通过指标找到模型熟悉的视角、恢复性能；多视角训练后，是否仍有更好的候选视角值得选择。')
add(p,'image',file='experiment_overview.png',width=16.5)
add(p,'caption','图 1  本轮对照实验。上方为同一任务初态的实际候选图像。')
table(p,['工作','目的与进展'],[
 ['训练覆盖对照','补训与 Broad64 设置匹配的规范视角模型，排除其他训练差异。模型已完成，正在比较两者对相同候选视角的适应能力。'],
 ['候选视角评测','整理多噪声下的候选结果，比较固定视角与逐起点 Oracle。旧恢复状态实验已整理，新一轮从标准初态评测完整任务。'],
 ['指标与信息性分析','绘制成功率、Accel、可见性的空间分布，比较单指标和组合规则；结合信息视角对照，区分训练位姿熟悉度与可见性增加的作用。'],
],[.24,.76])
add(p,'p','评测中发现并修复了渲染重复性问题，现已统一候选打分与闭环输入。后续用修复后的数据重新检查已有发现。')

p=page()
add(p,'h1','二、结果与判断')
add(p,'p','以下为已有 Broad64 实验：8 个任务、64 个恢复起点，包含任务中间状态。新的标准初态整任务评测仍在进行。')
add(p,'image',file='historical_spatial_maps.png',width=15.8)
add(p,'caption','图 2  候选空间的成功率、平均 Accel 与可见性，颜色为各起点均值；坐标为方位角、仰角偏移及半径倍率，星号为规范视角。')
table(p,['视角选择','续跑成功率','目前的判断'],[
 ['规范视角','64.70%','共同参照'],
 ['均匀选择 / 最低 Accel','58.31% / 61.38%','Accel 高于均匀选择，但未超过规范视角'],
 ['Accel Top-10 + 可见性','60.50%','组合后未进一步改善'],
 ['逐起点经验 Oracle','75.88%','高于规范视角 11.18 个百分点，仍有待识别的候选收益'],
],[.36,.22,.42])
add(p,'p','Oracle 按每个起点的多噪声平均成功率选出最好视角。当前候选收益尚未被指标识别，其稳定性还需修复后复评。')
add(p,'p','此前可见性对照中，信息视角为 75.0%，规范视角为 83.3%（两个任务、12 个恢复状态），可见像素增加未带来成功率提升。')
add(p,'h1','三、下一步')
add(p,'p','1. 重画两模型的热力图和差值图，分任务、初态对照成功率、Accel 及其噪声波动，检查高成功率区域与训练视角的距离。')
add(p,'p','2. 同时检验两类收益：相对保持扰动视角，指标选择能恢复多少性能；相对规范及每任务固定视角，逐初态 Oracle 还剩多少空间。')
add(p,'p','3. 根据分布分析设计组合指标或轻量选择器，在留出初态上比较实际收益与 Oracle 差距，优先复用当前评测数据。')
add(p,'p','4. 逐步开展 RoboCasa 小规模试验，先选少量任务跑通数据和评测流程，再加入未见场景对照，检查 LIBERO 中的指标规律能否迁移。')
add(p,'p','5. 同步准备真机实验，先完成相机标定和少量候选位姿设置，对比扰动视角、规范视角和指标所选视角；之后再评估实际移动与在线搜索。')

def esc(x):return html.escape(str(x)).replace('\n','<br/>')
def html_block(b):
    k=b['kind'];t=esc(b.get('text',''))
    if k=='meta':t=t.replace('26B951131','<span style="font-family:latin">26B951131</span>')
    if k=='table':
        ws=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
        head='<tr>'+''.join(f'<th style="width:{w*100:.1f}%">{esc(v)}</th>' for v,w in zip(b['heads'],ws))+'</tr>'
        rows=''.join('<tr>'+''.join(f'<td>{esc(v)}</td>' for v in row)+'</tr>' for row in b['rows'])
        return '<table>'+head+rows+'</table>'
    if k=='image':return f'<div class="image" style="text-align:center"><img src="{b["file"]}" width="{b["width"]/17.2*100:.1f}%"/></div>'
    if k in ['h1','h2','h3']:return f'<{k}>{t}</{k}>'
    return f'<p class="{k}">{t}</p>'

css='''
@font-face {font-family:report;src:url(NotoSansCJK-Regular.ttc);}
@font-face {font-family:report;src:url(NotoSansCJK-Bold.ttc);font-weight:bold;}
@font-face {font-family:latin;src:url(DejaVuSans.ttf);}
body {font-family:report;font-size:10.5pt;line-height:1.48;color:#171717;}
p {margin:0 0 5pt 0;} .title{font-size:20pt;font-weight:bold;margin-bottom:13pt;}
.meta {font-size:10pt;margin-bottom:3pt;} h1{font-size:15pt;margin:13pt 0 8pt;}
h2{font-size:11.5pt;margin:10pt 0 6pt;} h3{font-size:11.5pt;margin:9pt 0 7pt;}
table{width:100%;border-collapse:collapse;margin:5pt 0 9pt;table-layout:fixed;}
th,td{border:0.55pt solid #cdd3d7;padding:5pt 6pt;vertical-align:top;font-size:9.6pt;line-height:1.4;}
th{background:#edf0f2;text-align:left;font-weight:bold;}
.caption{font-size:9pt;color:#555;line-height:1.4;margin:2pt 0 9pt;}
.note{font-size:10pt;background:#f2f3f4;padding:8pt;line-height:1.45;margin-top:7pt;}
.small{font-size:8.5pt;color:#666;margin-top:9pt;}.image{margin:4pt 0;}
'''
pdf=fitz.open();archive=fitz.Archive([str(ASSETS),'/usr/share/fonts/opentype/noto','/usr/share/fonts/truetype/dejavu'])
qa=[]
for i,blocks in enumerate(pages):
    pg=pdf.new_page(width=595.28,height=841.89)
    body=''.join(html_block(b) for b in blocks)
    spare,scale=pg.insert_htmlbox(fitz.Rect(54,46,541.28,791),body,css=css,archive=archive,scale_low=1)
    if spare < 0:raise RuntimeError(f'Page {i+1} overflow; shorten content, do not auto shrink')
    pg.insert_htmlbox(fitz.Rect(54,806,470,830),'科研进展周报 · VLA 视角利用 · 2026.08.27—09.09',css=css+'body{font-size:8pt;color:#777;}',archive=archive)
    pg.insert_text((527,818),str(i+1),fontsize=8,color=(.45,.45,.45))
    qa.append({'page':i+1,'spare_height':spare,'scale':scale})
pdf.set_metadata({'title':STEM,'author':'龙君昱','subject':'VLA 视角适应性与评价指标研究'})
pdf.save(OUT/(STEM+'.pdf'),garbage=4,deflate=True)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(4)
for name,size in [('Title',20),('Heading 1',15),('Heading 2',11.5),('Heading 3',11.5)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True
    st.font.color.rgb=RGBColor.from_string('171717');st.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0);st.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0]
footer.text='科研进展周报 · VLA 视角利用\t'
footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.8),WD_ALIGN_PARAGRAPH.RIGHT)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('777777')
def fmt_runs(par,size=10.5,bold=False,color='171717'):
    for r in par.runs:
        r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')

for pi,blocks in enumerate(pages):
    if pi:doc.add_page_break()
    for b in blocks:
        k=b['kind'];text=b.get('text','')
        if k=='table':
            tb=doc.add_table(rows=1,cols=len(b['heads']));tb.style='Table Grid';tb.autofit=False
            widths=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
            for c,w in zip(tb.columns,widths):c.width=Cm(17.2*w)
            for rownum,values in enumerate([b['heads']]+b['rows']):
                row=tb.rows[0] if rownum==0 else tb.add_row()
                cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
                if rownum==0:
                    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                for cell,v,w in zip(row.cells,values,widths):
                    cell.width=Cm(17.2*w);cell.text=v
                    if rownum==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');cell._tc.get_or_add_tcPr().append(sh)
                    for par in cell.paragraphs:
                        par.paragraph_format.space_after=Pt(4);par.paragraph_format.space_before=Pt(3)
                        par.paragraph_format.line_spacing=1.2;fmt_runs(par,9.5,rownum==0)
            continue
        if k=='image':
            par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after=Pt(2);par.add_run().add_picture(str(ASSETS/b['file']),width=Cm(b['width']))
            continue
        style={'title':'Title','h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(k)
        par=doc.add_paragraph(text,style=style)
        if k in ['caption','small','meta']:
            fmt_runs(par,9 if k!='meta' else 10,color='555555' if k!='meta' else '171717')
            par.paragraph_format.line_spacing=1.2;par.paragraph_format.space_after=Pt(4)
        if k=='note':
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F3F4');par._p.get_or_add_pPr().append(sh)
            fmt_runs(par,10);par.paragraph_format.line_spacing=1.25
doc.core_properties.title=STEM;doc.core_properties.author='龙君昱';doc.core_properties.subject='VLA 视角利用研究双周进展'
doc.save(OUT/(STEM+'.docx'))

# Audit text coverage and package integrity; PDF is rendered from the same blocks,
# not an assertion of Word's platform-dependent pagination.
import zipfile
with zipfile.ZipFile(OUT/(STEM+'.docx')) as z:assert z.testzip() is None
import unicodedata
texts=unicodedata.normalize('NFKC','\n'.join(pg.get_text() for pg in pdf))
for required in ['龙君昱','26B951131','64.70','75.88','一、主要工作','三、下一步','扰动','熟悉','渲染']:
    assert required in texts, required
assert len(pdf)==2
assert '事务型工作' not in texts
for i,pg in enumerate(pdf):pg.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(ASSETS/f'page-{i+1}.png')
from PIL import Image,ImageOps,ImageDraw
thumbs=[]
for i in range(len(pdf)):
    im=Image.open(ASSETS/f'page-{i+1}.png').convert('RGB');im.thumbnail((358,506));thumbs.append(im)
contact=Image.new('RGB',(3*378,536),'#ddd')
for i,im in enumerate(thumbs):contact.paste(im,((i%3)*378+10,(i//3)*536+10))
contact.save(ASSETS/'contact.png')
manifest={'period':['2026-08-27','2026-09-09'],'report_date':'2026-09-09','revision_date':'2026-09-10',
 'sources':[{'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sources],
 'progress_snapshot':snapshot,'historical_plot_values':dict(zip(keys,vals)),
 'pdf_layout':qa,'docx_checks':['zip integrity','shared content blocks','explicit 2-page sections'],
 'pdf_note':'Same-content typeset reading copy; not a LibreOffice conversion of the DOCX.',
 'new_experiments_started':False,'heldout_outcomes_read':False,
 'outputs':{ext:hashlib.sha256((OUT/(STEM+'.'+ext)).read_bytes()).hexdigest() for ext in ['docx','pdf']}}
(OUT/'_provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'outputs':[str(OUT/(STEM+'.'+x)) for x in ['docx','pdf']], 'layout':qa},ensure_ascii=False,indent=2))
