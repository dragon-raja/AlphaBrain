"""Read-only research evidence -> editable biweekly report and same-content PDF.

No policy execution, training, or evaluation-controller modification.
Dependencies are isolated in /tmp, not installed into the experiment environment.
"""
from pathlib import Path
import sys, json, hashlib, html
sys.path[:0] = ['/tmp/dsol-weekly-docx-deps-20260909',
                '/root/.cache/uv/archive-v0/XO8_y4zBEt33DOY7wfN64']
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, fontManager
import pymupdf as fitz
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent
ASSETS = OUT / '_assets'
ASSETS.mkdir(exist_ok=True)
REPO = Path('/workspace/projects/alphabrain-dsol-paper1')
EXP = Path('/share/longjunyu/alphabrain/experiments')
NEW = EXP / 'dsol-standard-initial-aa0-v1-20260908'
OLD = EXP / 'dsol-view-landscape-v1-20260908'
STEM = '科研进展双周报_20260827-20260909'
font = FontProperties(fname='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
plt.rcParams.update({'font.family': font.get_name(), 'axes.unicode_minus': False,
                     'font.size': 11, 'axes.spines.top': False, 'axes.spines.right': False})

sources = [
    REPO/'docs/dsol_paper1/training_match_execution_status_20260907_zh.md',
    REPO/'docs/dsol_paper1/libero_expanded_a_constructed_gate_20260827_zh.md',
    REPO/'docs/dsol_paper1/view_repeatability_root_cause_20260908_zh.md',
    REPO/'docs/dsol_paper1/render_protocol_bridge_and_historical_validity_20260908_zh.md',
    REPO/'docs/dsol_paper1/standard_initial_aa0_execution_20260908_zh.md',
    OLD/'checkpoint-audit/completed_run_receipt.json',
    OLD/'metric-oracle-comparison-v2/broad_summary.json',
    OLD/'render-protocol-bridge-v1/analysis-recovered-v1.json',
    NEW/'release.json', NEW/'assets/aa0-a/00/policy_inputs.npz',
]

rules = json.loads(sources[6].read_text())['overall']['rules']
keys = ['canonical', 'min_accel', 'accel_top10_max_visibility', 'statewise_empirical_max_O32']
vals = [rules[k]['success']['mean']*100 for k in keys]
assert np.allclose(vals, [64.697265625,61.376953125,60.498046875,75.87890625])
fig,ax = plt.subplots(figsize=(7.2,2.7),layout='constrained')
colors = ['#8C99A3','#486C8E','#7098A8','#B3A085']
ax.bar(range(4),vals,color=colors,width=.56)
for i,v in enumerate(vals): ax.text(i,v+2,f'{v:.2f}%',ha='center',fontsize=12)
ax.set_xticks(range(4),['规范视角','最低平均 Accel','Accel 前10名\n再按可见性选取','逐恢复起点\n经验 Oracle'])
ax.set_ylim(0,100);ax.set_ylabel('恢复后续跑成功率（%）')
ax.set_yticks([0,25,50,75,100]);ax.grid(axis='y',alpha=.18);ax.set_axisbelow(True)
fig.savefig(ASSETS/'historical_metrics.png',dpi=180);plt.close(fig)

with np.load(sources[9]) as a:
    ids = a['candidate_ids'].tolist()
    chosen = ['canonical','broad_train_000','broad_heldout_000']
    fig,axs=plt.subplots(1,3,figsize=(7.2,2.4),layout='constrained')
    for ax,c,label in zip(axs,chosen,['规范视角','训练目录候选 000','留出目录候选 000']):
        ax.imshow(a['external_images'][ids.index(c)]);ax.set_title(label,fontsize=11);ax.axis('off')
    fig.savefig(ASSETS/'official_initial_views.png',dpi=180);plt.close(fig)

# Freeze a progress snapshot; intentionally do not read held-out outcome labels.
snapshot_file = OUT/'_progress_snapshot.json'
if snapshot_file.exists():
    snapshot=json.loads(snapshot_file.read_text())
else:
    from datetime import datetime, timezone
    snapshot={'observed_at_utc':datetime.now(timezone.utc).isoformat(),'hosts':{}}
    for host in ['fresh','gnu']:
        d=NEW/'hosts'/host
        snapshot['hosts'][host]={
            'status':json.loads((d/'status.json').read_text()),
            'first_wave':json.loads((d/'dense-init-0-complete.json').read_text()),
            'initial_gate':json.loads((d/'gate-pass.json').read_text()),
            'dynamic_gate':json.loads((NEW/'scheduling/dynamic-v1'/host/'cross-host-gate-pass.json').read_text()),
        }
    snapshot_file.write_text(json.dumps(snapshot,ensure_ascii=False,indent=2)+'\n')
assert all(v['first_wave']['episodes']==24832 for v in snapshot['hosts'].values())
from datetime import datetime, timedelta
asof = (datetime.fromisoformat(snapshot['observed_at_utc'])+timedelta(hours=8)).strftime('%m 月 %d 日 %H:%M')

pages=[]
def page():
    x=[];pages.append(x);return x
def add(p,kind,text='',**kw):p.append(dict(kind=kind,text=text,**kw))
def table(p,heads,rows,widths=None):add(p,'table',heads=heads,rows=rows,widths=widths)

p=page()
add(p,'title','科研进展周报')
add(p,'meta','汇报人：____________    学号：____________')
add(p,'meta','研究方向：视觉—语言—动作模型（VLA）的视角利用')
add(p,'meta','统计周期：2026 年 8 月 27 日—9 月 9 日（双周）')
add(p,'meta',f'汇报日期：2026 年 9 月 9 日    进度截点：北京时间 {asof}')
add(p,'h1','一、研究进展及成果汇报')
for s in [
 '1. 明确当前主问题：在多视角后训练后，从任务初态选择一个固定外部视角，是否仍能提高整任务成功率，以及这种收益能否由指标或学习器识别。',
 '2. 完成规范视角匹配后训练，使其与已有 Broad64 模型使用相同初始化、数据记录和优化预算，为分析训练覆盖如何改变视角价值建立对照。',
 '3. 完成旧候选矩阵与可见性实验整理。恢复起点上的经验最佳视角优于若干指标规则，但可见像素增加并不自动带来成功率提升；上述结果仍需在新主协议下验证。',
 '4. 定位并修复闭环评测中的渲染非确定性，完成新旧协议配对复核；建立候选打分图像与实际策略输入的一致性检查。',
 '5. 启动双机标准初态整任务评测，预算为 198,656 次；两机已共同完成第一轮 49,664 次，其余轮次继续执行。当前不报告未完成矩阵的总体成功率。',
]:add(p,'p',s)
add(p,'h1','二、本阶段的研究进展')
add(p,'h2','科研型工作')
add(p,'h3','1. 完成规范视角与多视角训练的匹配对照')
add(p,'p','前期多视角后训练已观察到相机扰动下的性能改善，但部分旧对照的训练配方并不完全一致。本阶段以 Broad64 M-B 为锚，补训规范视角模型，避免将继续训练预算差异误认为视角覆盖的效果。')
table(p,['控制项','本阶段完成情况'],[
 ['共同训练条件','同一 π0.5 初始权重、同一数据容器及记录划分；保留相同腕部图、语言与动作监督。'],
 ['主要实验变量','外部训练图像：规范位姿图像 / Broad64 多位姿图像。'],
 ['优化预算','seed 41；全局 batch 32；2,000 次更新；64,000 次样本呈现；学习率轨迹匹配。'],
 ['完成与边界','9 月 7 日训练结束，保存配置及权重身份审计通过；当前仅一组匹配训练 seed，闭环对照正在执行。'],
],[.22,.78])

p=page()
add(p,'h3','2. 信息视角与候选价值：已有结果及其适用范围')
add(p,'p','本阶段对可见性构造、候选穷举与指标选择进行了联合复核。重点从“某个视角曾经成功”转向“同一候选在多套推理噪声下的平均表现”，并区分任务差异、初始条件差异与噪声波动。')
add(p,'h3','（1）可见性增加不等同于任务信息有效增加')
add(p,'p','8 月 27 日的构造实验在两个任务、12 个测试恢复状态上比较规范、信息视角和等幅移动对照。旧协议下，三者续跑成功率分别为 83.3%、75.0% 和 66.7%。信息视角相对等幅对照有方向性收益，但未超过规范视角，样本区间也较宽。')
add(p,'p','随后匹配的视角支持补训中，信息视角条件仍为 75.0%，未形成进一步收益。这提示实体像素占比只能作为候选特征；现有结果不足以判定具体失败机制，也不能据此宣称外部视角或视觉信息无用。')
add(p,'h3','（2）旧候选矩阵：经验空间与指标能力存在差距')
add(p,'p','完成 Broad64 模型在 8 个任务、64 个恢复起点、97 个候选、每候选 32 套噪声上的旧矩阵整理。各候选先对噪声求平均，再比较选择规则；下图为同一总体上的描述性结果。')
add(p,'image',file='historical_metrics.png',width=16.4)
add(p,'caption','图 1  旧渲染协议下的恢复后续跑结果。Accel 指 Flow 去噪轨迹的加速度指标，而非机械臂运动加速度；组合规则沿用旧可见性口径。柱形为点估计，不代表独立确认收益。')
add(p,'p','经验 Oracle 为 75.88%，规范视角为 64.70%，差 11.18 个百分点；最低平均 Accel 与组合规则未超过规范视角。这表明该旧矩阵中的事后候选空间尚未被这些规则有效识别。')
add(p,'note','结论边界：这些起点包含演示中间状态及构造场景；75.88% 不是标准初态整任务成绩，也不是任务中途动态选视角的成绩。经验最大值还受有限噪声选择偏差影响，旧闭环另含渲染随机性。因此本图仅作为后续研究动机，不作为新协议下已确认的主结论。')

p=page()
add(p,'h3','3. 定位评测非确定性并完成协议修复')
add(p,'p','在重复评测核查中发现：即使 Flow 噪声和任务起点相同，少量图像像素仍可能改变，随后引发动作及物理轨迹分叉。通过相同请求重复推理、固定动作回放和静态重复渲染，最早分歧被定位到 MuJoCo/EGL 多重采样抗锯齿路径，而不是已对齐的 Flow 初噪。')
table(p,['检查或修复','实测结果与解释'],[
 ['相同策略请求重复推理','原诊断的 790 次比较中，动作逐位相同；支持这些请求上的推理可重复。'],
 ['固定动作回放与静态渲染','固定动作下物理轨迹一致；旧抗锯齿渲染可出现少量通道值相差 1，闭环中已观察到成败翻转。'],
 ['关闭多重采样（AA0）','显式设置并检查实际 framebuffer 样本数为 0；不改变模型、Flow 去噪或控制步长。'],
 ['1,536 次新旧协议配对复核','AA0 下两模型各 192 对完整轨迹一致；AA4 下轨迹差异分别为 35/192、20/192，成败翻转为 0/192、2/192。'],
 ['新初态及双机验收','初始策略输入与候选指标图逐项一致；跨机重复与后续动态调度验收均通过后才继续正式评测。'],
],[.29,.71])
add(p,'p','关闭抗锯齿也会改变图像分布，因此不能把它仅视为消除无害抖动。桥接实验中，两模型性能差距的新旧变化为 +0.52 个百分点，95% 区间为 [−5.73, 7.29]，尚不足以确认预设 ±5 个百分点范围内的等价性。')
add(p,'note','处理原则：保留历史数据与结论来源，但不将其重新认证为新协议结果；后续候选图、指标和闭环标签统一使用 AA0。该问题主要涉及在线评测，当前没有因此重训全部模型的依据。')
add(p,'h3','4. 统一研究定义，避免不同成功率口径混用')
table(p,['对象','本阶段调整'],[
 ['主实验起点','由演示恢复快照改为官方标准任务初态，完整执行至成功或统一超时。'],
 ['视角选择方式','每初态选一次外部位姿，整条 rollout 固定；腕部相机随机器人运动不等于外部相机重新选择。'],
 ['经验 Oracle','每初态、每视角先平均 32 套噪声，再在候选中取最大；不同初态可选不同视角，不逐噪声挑赢家。'],
 ['暂缓内容','中间状态动态选视角、主动移动相机与腕部动作耦合；既有恢复实验仅保留为辅助分析。'],
],[.23,.77])

p=page()
add(p,'h3','5. 双机标准初态整任务评测已启动')
table(p,['设置','当前冻结内容'],[
 ['任务与初态','原候选研究的 8 个已知任务；每任务 4 个官方初态，共 32 个。不是 LIBERO 全套件。'],
 ['候选与模型','97 个视角：规范 1 + 训练目录 64 + 留出目录 32；规范匹配模型 / Broad64 模型。'],
 ['推理与输入','每初态 32 套显式噪声，跨模型、视角按重规划索引对齐；外部图 + 腕部图 + 任务语言。'],
 ['科学预算','2 × 32 × 97 × 32 = 198,656 次完整任务执行；4 轮，每轮 49,664 次。'],
 ['开发与测试','初态 0、1 用于开发，2、3 留出；不是未见任务评测，亦未证明与 VLA 训练来源完全隔离。'],
 ['双机执行','fresh、gnu 各 8 卡，各负责 16 个初态并运行两个模型；每机 32 个仿真 worker。'],
],[.22,.78])
add(p,'image',file='official_initial_views.png',width=16.4)
add(p,'caption','图 2  新协议实际策略输入示例：奶酪入碗任务、官方初态 0。按目录固定展示三个候选，未依据成功率选择；三图共享同一物理初态。新主实验不增加遮挡板或背景变化。')
add(p,'p',f'截至北京时间 {asof}，第一轮两个模型、全部 8 个任务的 49,664 次评测已完整完成；两机均处于第二轮 Broad64 评测。首轮完成不代表初态覆盖充分，本周报不读取留出标签来挑选方法，也不提前汇总不平衡的总体成功率。')
add(p,'p','9 月 9 日将每机内部的静态 worker 分片改为共享任务队列，以减少长短回合导致的尾部空闲。优化未改变样本集合、并发上限及模型推理协议，额外跨机、跨 GPU、重复轨迹验收通过后恢复执行。')
add(p,'note','当前执行预算仅含每任务 4 个初态；后续拟补至每任务 8 个，需另行冻结剩余初态及预算，不自动追加。扩展时关注初态间差异与选择器数据需求，不会因收益方向而挑选追加样本。')

p=page()
add(p,'h2','事务型工作')
add(p,'p','1. 统一研究进展材料入口，区分历史恢复实验与新初态主实验；保留旧报告、协议和数据来源，减少重复文件造成的口径混淆。')
add(p,'p','2. 整理模型匹配审计、噪声库、候选图像和闭环记录的关联；新旧协议分目录归档。建立双机独立写入与完整性验收，避免重复执行和结果混入。')
add(p,'h1','三、下阶段工作安排')
add(p,'h2','科研型工作')
table(p,['优先级','工作内容','预期产出'],[
 ['1','完成已冻结的 4 初态整任务矩阵，检查缺失、重复及输入一致性。','两模型同口径的完整结果表；逐任务、逐初态差异及不确定性。'],
 ['2','比较规范、全局最佳固定、每任务最佳固定、每初态经验 Oracle。','明确多视角训练改变了哪些视角区域，以及剩余可选收益有多大。'],
 ['3','绘制实际离散候选的成功率 / Accel 空间切片及两模型差值；比较可见性与 Accel Top-10 组合。','检验指标与闭环价值是否相关，哪些任务有效、哪些任务失效。'],
 ['4','在开发初态上训练轻量价值回归或排序选择器；将规范视角保留为候选，固定方法后测试。','与指标规则同输入权限的留出初态成功率；必要时对少数最终方法补独立噪声复评。'],
 ['5','依据完整结果评估样本覆盖和论文主张，再决定扩展初态或外部验证。','优先补关键证据；RoboCasa、真机和主动相机暂不并行扩展。'],
],[.10,.48,.42])
add(p,'p','上述指标比较和静态选择器测试可复用矩阵中相应候选的真实闭环标签，不需为每条规则重跑全候选。新增图像特征可能需要补算；新增初态、候选位姿或动态相机策略则需要新闭环。')
add(p,'h2','需重点关注的问题')
add(p,'p','• 样本独立性：本轮只有 32 个任务初态，不能把候选与噪声重复当成更多独立场景；选择器结论先限定为已知任务内的初态留出。')
add(p,'p','• 观察权限：允许读取全部候选图和 Accel 的选择器，是候选图可查询条件，不代表静止实体相机能够免费获得其他视角。')
add(p,'p','• 论文贡献：当前目标是建立“训练覆盖如何改变视角价值、指标能利用多少剩余收益”的证据链；尚未证明稳定正向选择方法，也不把动态主动感知列为已完成贡献。')
add(p,'h2','事务型工作')
add(p,'p','完成主矩阵后更新统一研究报告，保留图表的任务范围、初态数量、噪声数量及渲染版本；与师兄讨论指标失效模式和后续最小验证方案。')
add(p,'small','资料依据：匹配训练完成回执（9/7–9/8）；构造可见性实验记录（8/27）；旧候选矩阵统计（9/8）；渲染协议配对复核（9/8）；标准初态冻结协议与双机进度（9/9）。本报告未新增训练或评测，配图直接取自归档统计与新协议输入。')

def esc(x):return html.escape(str(x)).replace('\n','<br/>')
def html_block(b):
    k=b['kind'];t=esc(b.get('text',''))
    if k=='table':
        ws=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
        head='<tr>'+''.join(f'<th style="width:{w*100:.1f}%">{esc(v)}</th>' for v,w in zip(b['heads'],ws))+'</tr>'
        rows=''.join('<tr>'+''.join(f'<td>{esc(v)}</td>' for v in row)+'</tr>' for row in b['rows'])
        return '<table>'+head+rows+'</table>'
    if k=='image':return f'<div class="image"><img src="{b["file"]}" width="100%"/></div>'
    if k in ['h1','h2','h3']:return f'<{k}>{t}</{k}>'
    return f'<p class="{k}">{t}</p>'

css='''
@font-face {font-family:report;src:url(NotoSansCJK-Regular.ttc);}
@font-face {font-family:report;src:url(NotoSansCJK-Bold.ttc);font-weight:bold;}
body {font-family:report;font-size:10.5pt;line-height:1.48;color:#171717;}
p {margin:0 0 7pt 0;} .title{font-size:23pt;font-weight:bold;margin-bottom:13pt;}
.meta {font-size:10pt;margin-bottom:3pt;} h1{font-size:15pt;margin:13pt 0 8pt;}
h2{font-size:11.5pt;margin:10pt 0 6pt;} h3{font-size:11.5pt;margin:9pt 0 7pt;}
table{width:100%;border-collapse:collapse;margin:5pt 0 9pt;table-layout:fixed;}
th,td{border:0.55pt solid #cdd3d7;padding:5pt 6pt;vertical-align:top;font-size:9.6pt;line-height:1.4;}
th{background:#edf0f2;text-align:left;font-weight:bold;}
.caption{font-size:9pt;color:#555;line-height:1.4;margin:2pt 0 9pt;}
.note{font-size:10pt;background:#f2f3f4;padding:8pt;line-height:1.45;margin-top:7pt;}
.small{font-size:8.5pt;color:#666;margin-top:9pt;}.image{margin:4pt 0;}
'''
pdf=fitz.open();archive=fitz.Archive([str(ASSETS),'/usr/share/fonts/opentype/noto'])
qa=[]
for i,blocks in enumerate(pages):
    pg=pdf.new_page(width=595.28,height=841.89)
    body=''.join(html_block(b) for b in blocks)
    spare,scale=pg.insert_htmlbox(fitz.Rect(54,46,541.28,791),body,css=css,archive=archive,scale_low=1)
    if spare < 0:raise RuntimeError(f'Page {i+1} overflow; shorten content, do not auto shrink')
    pg.insert_htmlbox(fitz.Rect(54,806,470,830),'科研进展周报 · VLA 视角利用 · 2026.08.27—09.09',css=css+'body{font-size:8pt;color:#777;}',archive=archive)
    pg.insert_text((527,818),str(i+1),fontsize=8,color=(.45,.45,.45))
    qa.append({'page':i+1,'spare_height':spare,'scale':scale})
pdf.set_metadata({'title':STEM,'author':'','subject':'双周科研进展；历史辅助证据与当前主评测分开呈现'})
pdf.save(OUT/(STEM+'.pdf'),garbage=4,deflate=True)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',23),('Heading 1',15),('Heading 2',11.5),('Heading 3',11.5)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True
    st.font.color.rgb=RGBColor.from_string('171717');st.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0);st.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0]
footer.text='科研进展周报 · VLA 视角利用\t'
footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.8),WD_ALIGN_PARAGRAPH.RIGHT)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('777777')
def fmt_runs(par,size=10.5,bold=False,color='171717'):
    for r in par.runs:
        r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')

for pi,blocks in enumerate(pages):
    if pi:doc.add_page_break()
    for b in blocks:
        k=b['kind'];text=b.get('text','')
        if k=='table':
            tb=doc.add_table(rows=1,cols=len(b['heads']));tb.style='Table Grid';tb.autofit=False
            widths=b.get('widths') or [1/len(b['heads'])]*len(b['heads'])
            for c,w in zip(tb.columns,widths):c.width=Cm(17.2*w)
            for rownum,values in enumerate([b['heads']]+b['rows']):
                row=tb.rows[0] if rownum==0 else tb.add_row()
                cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
                if rownum==0:
                    repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                for cell,v,w in zip(row.cells,values,widths):
                    cell.width=Cm(17.2*w);cell.text=v
                    if rownum==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');cell._tc.get_or_add_tcPr().append(sh)
                    for par in cell.paragraphs:
                        par.paragraph_format.space_after=Pt(4);par.paragraph_format.space_before=Pt(3)
                        par.paragraph_format.line_spacing=1.2;fmt_runs(par,9.5,rownum==0)
            continue
        if k=='image':
            par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after=Pt(2);par.add_run().add_picture(str(ASSETS/b['file']),width=Cm(b['width']))
            continue
        style={'title':'Title','h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(k)
        par=doc.add_paragraph(text,style=style)
        if k in ['caption','small','meta']:
            fmt_runs(par,9 if k!='meta' else 10,color='555555' if k!='meta' else '171717')
            par.paragraph_format.line_spacing=1.2;par.paragraph_format.space_after=Pt(4)
        if k=='note':
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F3F4');par._p.get_or_add_pPr().append(sh)
            fmt_runs(par,10);par.paragraph_format.line_spacing=1.25
doc.core_properties.title=STEM;doc.core_properties.author='';doc.core_properties.subject='VLA 视角利用研究双周进展'
doc.save(OUT/(STEM+'.docx'))

# Audit text coverage and package integrity; PDF is rendered from the same blocks,
# not an assertion of Word's platform-dependent pagination.
import zipfile
with zipfile.ZipFile(OUT/(STEM+'.docx')) as z:assert z.testzip() is None
texts='\n'.join(pg.get_text() for pg in pdf)
for required in ['198,656','49,664','75.88','64.70','一、研究进展','三、下阶段','渲染']:
    assert required in texts, required
assert len(pdf)==5
for i,pg in enumerate(pdf):pg.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(ASSETS/f'page-{i+1}.png')
from PIL import Image,ImageOps,ImageDraw
thumbs=[]
for i in range(5):
    im=Image.open(ASSETS/f'page-{i+1}.png').convert('RGB');im.thumbnail((358,506));thumbs.append(im)
contact=Image.new('RGB',(3*378,2*536),'#ddd')
for i,im in enumerate(thumbs):contact.paste(im,((i%3)*378+10,(i//3)*536+10))
contact.save(ASSETS/'contact.png')
manifest={'period':['2026-08-27','2026-09-09'],'report_date':'2026-09-09',
 'sources':[{'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sources],
 'progress_snapshot':snapshot,'historical_plot_values':dict(zip(keys,vals)),
 'pdf_layout':qa,'docx_checks':['zip integrity','shared content blocks','explicit 5-page sections'],
 'pdf_note':'Same-content typeset reading copy; not a LibreOffice conversion of the DOCX.',
 'new_experiments_started':False,'heldout_outcomes_read':False,
 'outputs':{ext:hashlib.sha256((OUT/(STEM+'.'+ext)).read_bytes()).hexdigest() for ext in ['docx','pdf']}}
(OUT/'_provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'outputs':[str(OUT/(STEM+'.'+x)) for x in ['docx','pdf']], 'layout':qa},ensure_ascii=False,indent=2))
