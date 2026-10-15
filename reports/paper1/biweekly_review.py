"""Evidence-backed A4 biweekly review; PDF/DOCX/Markdown share content blocks."""
import argparse
import html
from pathlib import Path

import numpy as np

from AlphaBrain.research.dsol.data.artifacts import read, sha, write
from AlphaBrain.research.dsol.plotting.window_comparison import curves, paired_stage_margin
from .confirmation_update import number, pp


def window_panel_legend(ax):
    """Keep all four line meanings visible without covering curve values."""
    handles, labels = ax.get_legend_handles_labels()
    labels = [label.replace('允许规范位', '允许选择规范位')
              .replace('不调整 ', '扰动不调整 ')
              .replace('规范 ', '规范视角 ') for label in labels]
    return ax.legend(handles, labels, loc='upper center', bbox_to_anchor=(.5, -.29),
                     ncol=2, fontsize=8.2, frameon=False, columnspacing=1.1)


def comparison_panels(summary):
    """Six panels; bridge uses matched four-rollout results, not the 32-rollout row."""
    panels = []
    for source, label in [('libero', 'LIBERO → 自建97 · 32次执行'),
                          ('plus', 'Plus → 官方候选 · 4次执行'),
                          ('bridge', 'Plus → 自建97 · 复用4次执行')]:
        for model, training in [('canonical', '规范训练'), ('broad', '多视角训练')]:
            if source == 'bridge':
                methods = summary['bridge'][model]
                if any(methods[m]['execution_repeats'] != 4 for m in ('include', 'exclude')):
                    raise ValueError('Bridge must use matched four-rollout estimates')
                values = {m: methods[m]['mean'] for m in ('include', 'exclude')}
            else:
                methods = summary['panels'][source+'_'+model]['methods']
                values = {m: methods[m]['scopes']['all']['mean'] for m in ('include', 'exclude')}
            panels.append(dict(title=label+'\n'+training, **values))
    return panels


def content(cfg, summary, transitions, audit):
    """Dates and report facts are separate from typography; no policy execution."""
    if cfg['evidence_cutoff'] > cfg['report_date']:
        raise ValueError('Evidence cannot come from after the report date')
    execution_trials=sum(m['execution_trials'] for m in audit['models'].values())
    pages=[[],[],[]]
    def add(page,kind,text='',**kwargs):pages[page].append(dict(kind=kind,text=text,**kwargs))
    def table(page,heads,rows,widths):add(page,'table',heads=heads,rows=rows,widths=widths)
    add(0,'title','科研进展双周报')
    add(0,'meta',f"汇报人：{cfg['author']}　　学号：{cfg['student_id']}")
    add(0,'meta','汇报周期：2026年10月1日—10月14日　｜　汇报日期：2026年10月15日')
    add(0,'h1','一、本期工作：新增模型对照与多步视角诊断')
    add(0,'p','围绕“如何选择视角以恢复模型性能”，本期新增XR-1模型对照，检验更好的模型表现能否带来更有效的Accel选位；同时将初态单次评分扩展为前1–5次重规划评分，完成RoboCasa独立验证及LIBERO、LIBERO-Plus对照。沿用已有权重，不重新训练模型。')
    table(0,['实验','完成内容'],[
        ['RoboCasa 独立确认',f"官方π0.5、XR-1各80个新初态／10个已知任务；97候选×8评分噪声。冻结各窗口终点后，用8套独立执行噪声确认；去重完成{execution_trials:,}次完整任务。"],
        ['LIBERO / Plus补充','同8任务×4初态、规范／多视角两模型。LIBERO复用每候选原32次闭环成绩；Plus完成10,912次官方相机扰动闭环，每格4次，另完成五窗口评分。'],
    ],[.23,.77])
    add(0,'caption','方法：各候选从同一初态短暂试执行，以前H次重规划的累计Accel选位，再复位执行完整任务。W1为初态评分，评分与执行噪声分离。正式执行中不切换候选；RoboCasa相机仍随底盘运动。')
    add(0,'h1','二、RoboCasa：规范成绩提高，选位收益仍不稳定')
    add(0,'image',file='robocasa_windows.png')
    add(0,'caption','图1　每模型80个新初态，seen／heldout各40个，划分布局与风格而非任务。蓝线为窗口1–5选位，橙线为扰动不调整，灰线为规范视角；完整任务评测使用8套新执行噪声。')
    gains=[]
    for m,label in [('official','π0.5'),('xr1','XR-1')]:
        g=summary['robocasa_original']['models'][m]['populations']['all']['methods']['window3']['gain_vs_stay']
        gains.append(f"{label}为{pp(g['mean'])} pp，95%区间[{pp(g['ci95'][0])}, {pp(g['ci95'][1])}]")
    add(0,'note','XR-1规范视角成功率为65.16%，高于π0.5的46.41%，但Accel选位未显示稳定收益。π0.5在W3、XR-1在W4/W5的实测均值均高于扰动不调整，说明存在改善迹象，但尚不能确认稳定恢复。')
    add(0,'p','π0.5：W1为39.53%，W3升至42.34%；XR-1：W1为52.03%，W4升至54.84%。多步评分能够改善部分选位，但收益没有随窗口长度稳定累积，需要进一步判断何时采用后续评分。')
    add(0,'caption','预先指定的W3相对扰动不调整：'+'；'.join(gains)+'。区间均包含零。W4/W5的改善为观察结果，尚不能据此确定通用窗口。')
    add(0,'h2','待核实：任务与评测设置的影响')
    add(0,'p','RoboCasa当前10个任务在规范视角下仍有较多失败；扰动使π0.5、XR-1分别下降5.42、13.52个百分点。恢复收益可能受任务能力、视角敏感性及相机设置影响。两环境的模型、输入与执行方式也不同，尚不能认定差异由benchmark本身造成；规范与扰动的差距也不是视角恢复的上限。')

    add(1,'h1','三、LIBERO与Plus：多步有效，但依赖候选空间')
    add(1,'image',file='libero_plus_windows.png')
    add(1,'caption','图2　三排均为同8任务、32初态。第三排以Plus扰动为不调整基线，复用LIBERO同初态的97候选评分及前4次闭环成绩，不能算作另一份独立验证；未评测从扰动位实际移动到目标位的过程。')
    add(1,'note','规范模型W1→W3：LIBERO为69.24%→85.16%；Plus→自建97为71.09%→87.50%，排除规范位后仍分别提高12.01、14.06个百分点。Plus官方候选中未见稳定的额外收益，多视角模型的增益也较小。多步诊断有作用，但更长不一定更好。')
    b=transitions['baseline_matching']
    add(1,'p',f"扰动基线不同：匹配4次执行时，自建扰动／Plus扰动的成功率，规范模型为{number(100*b['canonical']['custom4'])}%／{number(100*b['canonical']['plus4'])}%，多视角模型为{number(100*b['broad']['custom4'])}%／{number(100*b['broad']['plus4'])}%。多步收益需结合候选范围和训练条件判断，不能笼统概括为“Plus上无效”。")

    add(2,'h1','四、为什么窗口加长后反而下降？')
    rows=[]
    for key,label in [('libero97_canonical','LIBERO规范'),('robocasa_official','RoboCasa π0.5'),('robocasa_xr1','RoboCasa XR-1')]:
        vals=[]
        for pair in ('W1_W3','W3_W5'):
            a=transitions['panels'][key]['include'][pair]['all'];c=a['counts']
            vals.append(f"{c['improved']} / {c['degraded']} / {c['equal']}（{pp(a['delta'])} pp）")
        rows.append([label,*vals])
    table(2,['条件','W1→W3：改善/退化/相等','W3→W5：改善/退化/相等'],rows,[.26,.37,.37])
    add(2,'caption','表1　按初态比较所选视角的完整任务成功率，允许选择规范位。LIBERO每候选32次执行、RoboCasa8次；数量表示实测改善或退化，不代表逐初态显著性检验。')
    add(2,'p','LIBERO规范模型的改善覆盖6类任务；W3→W5的净下降约82%来自“上层抽屉放碗”。RoboCasa中改善与退化并存，XR-1在W3之后回升，当前结果不支持统一采用某个窗口。')
    add(2,'image',file='stable_misselection.png')
    add(2,'caption','图3　正柱表示新候选Accel更低。后两阶段评分使方法改选较差视角：LIBERO成功次数29/32→3/32，XR-1为8/8→2/8。两例的8套评分噪声均偏向较差候选，前后两半执行结果也均退化，提示问题不只是评分波动。')
    add(2,'note','已确认：累加后续评分有时会改选低成功率视角，评分一致并不保证选择正确。后续评分来自各候选执行后的不同状态，不是对初态的重复测量。待验证：这种误选是否与停滞、遮挡或任务进度差异有关。')
    add(2,'small','窗口时刻并不跨本体对齐：LIBERO／Plus为0、5、10、15、20动作步；RoboCasa π0.5为0、25、50、75、100；XR-1为0、16、32、48、64。H是重规划次数，不是相同物理时长或任务进度。')
    add(2,'h1','五、下一步工作')
    add(2,'p','① 分析适用条件：沿用预定的任务能力与视角敏感性分组，比较恢复收益；结合已有轨迹核查任务进度、动作推进与评分翻转的关系。\n② 改进选位规则：据此研究何时保留早期选择、何时采用后续评分，与W1及固定窗口公平比较；规则冻结后再独立验证。\n③ 验证实际执行：选位规则稳定后，研究少量查询下的相机搜索，再开展头部主动相机验证。当前候选试执行与复位仍是分析实验，不是可直接部署的方法。')
    add(2,'h2','实验质量与工程保障')
    add(2,'p','完成三机并行评测与结果去重；修复LIBERO重复复位造成的场景位置漂移，核验首帧及W1一致性。结果、源码与分析记录已归档共享盘。')
    return pages


def build(config_file):
    cfg=read(config_file); out=Path(cfg['output']).resolve()
    if out.exists():raise FileExistsError('Do not overwrite a biweekly snapshot')
    summary=read(cfg['synthesis']);transitions=read(cfg['transitions']);audit=read(cfg['robocasa_audit'])
    if any(s['status']!='PASS' for s in (summary,transitions,audit)):raise ValueError('Unaudited source')
    if summary['sources']!=transitions['sources']:raise ValueError('Evidence versions differ')
    if sum(m['execution_trials'] for m in audit['models'].values())!=44376:
        raise ValueError('Unexpected confirmation population')
    pages=content(cfg,summary,transitions,audit)
    assets=out/'_assets';assets.mkdir(parents=True)
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    import pymupdf as fitz
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    fontdir=Path(cfg['font_dir'])
    font_manager.fontManager.addfont(str(fontdir/'NotoSansCJK-Regular.ttc'))
    font=font_manager.FontProperties(fname=str(fontdir/'NotoSansCJK-Regular.ttc'))
    plt.rcParams.update({'font.family':font.get_name(),'axes.unicode_minus':False,'font.size':9})
    fig,axes=plt.subplots(1,2,figsize=(8.6,3.0))
    for ax,model,label in zip(axes,('official','xr1'),('π0.5 官方','XR-1 官方')):
        values=summary['panels']['robocasa_'+model]['methods']['include']['scopes']['all']['mean']
        curves(ax,values,title=label,limits=(20,80));ax.legend(fontsize=6.3,loc='upper left',ncol=2)
    fig.tight_layout();fig.savefig(assets/'robocasa_windows.png',dpi=220);plt.close(fig)
    fig,axes=plt.subplots(3,2,figsize=(8.6,8.6))
    for ax,panel in zip(axes.flat,comparison_panels(summary)):
        curves(ax,panel['include'],panel['exclude'],panel['title'])
        ax.set_title(panel['title'],fontsize=10.5,pad=10)
        window_panel_legend(ax)
    fig.tight_layout();fig.savefig(assets/'libero_plus_windows.png',dpi=220);plt.close(fig)
    rows=read(cfg['transition_rows']);fig,axes=plt.subplots(1,2,figsize=(8.6,2.8))
    for ax,(key,initial,title) in zip(axes,[('libero97_canonical',5,'LIBERO：上层抽屉放碗'),
                    ('robocasa_xr1','39dd9788d2b8a223cfc09850','XR-1：台面物体入柜')]):
        row=next(r for r in rows if r['panel']==key and r['initial']==initial and r['start']==3 and r['end']==5 and r['mode']=='include')
        paired_stage_margin(ax,row,title)
    fig.tight_layout();fig.savefig(assets/'stable_misselection.png',dpi=220);plt.close(fig)
    def esc(s):return html.escape(str(s)).replace('\n','<br/>')
    def block(b):
        kind=b['kind']
        if kind=='image':return f'<img src="{b["file"]}" style="width:100%">'
        if kind=='table':
            head=''.join(f'<th style="width:{100*w}%">{esc(t)}</th>' for t,w in zip(b['heads'],b['widths']))
            rows=''.join('<tr>'+''.join(f'<td style="width:{100*w}%">{esc(v)}</td>' for v,w in zip(r,b['widths']))+'</tr>' for r in b['rows'])
            return '<table><tr>'+head+'</tr>'+rows+'</table>'
        if kind in ('h1','h2'):return f'<{kind}>{esc(b["text"])}</{kind}>'
        return f'<p class="{kind}">{esc(b["text"])}</p>'
    css='''@font-face{font-family:cn;src:url(NotoSansCJK-Regular.ttc)}
    @font-face{font-family:cn;src:url(NotoSansCJK-Bold.ttc);font-weight:bold}
    body{font-family:cn;font-size:10.8pt;line-height:1.42;color:#1d252c}
    p{margin:0 0 6pt}.title{font-size:20pt;font-weight:bold;margin-bottom:11pt}
    .meta{font-size:9.2pt;color:#46515b;margin-bottom:3pt}
    h1{font-size:13.2pt;margin:11pt 0 7pt;color:#183b51}h2{font-size:10.5pt;margin:8pt 0 5pt;color:#183b51}
    table{width:100%;border-collapse:collapse;table-layout:fixed;margin:3pt 0 6pt}
    th,td{border:.5pt solid #cbd3d9;padding:4pt 5pt;vertical-align:top;font-size:9.2pt;line-height:1.35}
    th{background:#edf2f5;text-align:left}.caption{font-size:8.4pt;line-height:1.35;color:#59636b;margin:3pt 0 7pt}
    .note{font-size:10.2pt;background:#edf3f6;padding:6pt;line-height:1.4;margin-top:6pt}
    .small{font-size:8.1pt;color:#667078;line-height:1.35;margin-top:5pt}'''
    archive=fitz.Archive([str(assets),str(fontdir)]);pdf=fitz.open();layout=[]
    for i,blocks in enumerate(pages):
        pg=pdf.new_page(width=595.28,height=841.89)
        spare,scale=pg.insert_htmlbox(fitz.Rect(51,41,544,792),''.join(map(block,blocks)),css=css,archive=archive,scale_low=1)
        if spare<0:raise ValueError(f'Page{i+1} overflow; edit rather than shrink')
        pg.insert_htmlbox(fitz.Rect(51,807,544,830),f'科研进展双周报 · VLA视角利用 · 2026.10.15　　　　　　　　　{i+1} / 3',css=css+'body{font-size:7.5pt;color:#737d86}',archive=archive)
        layout.append(dict(page=i+1,spare_height=spare,scale=scale))
    stem='科研进展双周报_20261015'
    pdf.set_metadata(dict(title=stem,author=cfg['author'],subject='多窗口视角诊断：独立验证、适用条件与决策问题'))
    pdf.save(out/(stem+'.pdf'),garbage=4,deflate=True)
    for i,pg in enumerate(pdf):pg.get_pixmap(matrix=fitz.Matrix(1.4,1.4)).save(assets/f'page-{i+1}.png')
    doc=Document();sec=doc.sections[0]
    sec.page_width,sec.page_height=Cm(21),Cm(29.7)
    sec.top_margin,sec.bottom_margin=Cm(1.45),Cm(1.7);sec.left_margin=sec.right_margin=Cm(1.8)
    for name,size in [('Normal',10.8),('Title',20),('Heading 1',13.2),('Heading 2',10.5)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')
        st.paragraph_format.line_spacing=1.2;st.paragraph_format.space_after=Pt(5)
    md=[]
    for i,blocks in enumerate(pages):
        if i:doc.add_page_break()
        for b in blocks:
            k=b['kind']
            if k=='image':
                doc.add_picture(str(assets/b['file']),width=Cm(17.1));md.append(f'![图表](_assets/{b["file"]})')
            elif k=='table':
                tb=doc.add_table(rows=1,cols=len(b['heads']));tb.style='Table Grid'
                for j,row in enumerate([b['heads'],*b['rows']]):
                    cells=tb.rows[0].cells if j==0 else tb.add_row().cells
                    for cell,t,w in zip(cells,row,b['widths']):
                        cell.text=t;cell.width=Cm(17.4*w)
                        for p in cell.paragraphs:
                            for run in p.runs:run.font.size=Pt(8.5);run.bold=j==0
                md+=['| '+' | '.join(b['heads'])+' |','|'+'---|'*len(b['heads'])]
                md+=['| '+' | '.join(row)+' |' for row in b['rows']]
            else:
                p=doc.add_paragraph(b['text'],style={'title':'Title','h1':'Heading 1','h2':'Heading 2'}.get(k))
                if k in ('meta','caption','small'):
                    for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('59636B')
                if k=='note':
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF3F6');p._p.get_or_add_pPr().append(sh)
                md.append({'title':'# ','h1':'## ','h2':'### '}.get(k,'')+b['text'])
            md.append('')
    sec.footer.paragraphs[0].text='科研进展双周报 · 2026.10.15'
    doc.core_properties.title=stem;doc.core_properties.author=cfg['author'];doc.save(out/(stem+'.docx'))
    (out/(stem+'.md')).write_text('\n'.join(md))
    sources=[config_file,cfg['reference_pdf'],cfg['synthesis'],cfg['transitions'],cfg['transition_rows'],cfg['robocasa_audit'],__file__]
    write(out/'_provenance.json',dict(report_date=cfg['report_date'],period=cfg['period'],evidence_cutoff=cfg['evidence_cutoff'],
         sources={str(p):sha(p) for p in sources},layout=layout,new_experiments=0,
         outputs={suffix:sha(out/(stem+suffix)) for suffix in ('.pdf','.docx','.md')},
         note='Reference formatting reused; not a new independent analysis. DOCX pagination may vary with fonts. Current unified research PDF unchanged.'))
    print(out);print(layout)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--config',required=True);args=p.parse_args();build(args.config)
